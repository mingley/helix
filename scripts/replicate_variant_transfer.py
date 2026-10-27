#!/usr/bin/env python3
"""Independent raw-XLSX/NumPy replay for the frozen Siraj transfer study.

No imports from the primary predictor, source binder or evaluator. Empirical
reading requires a separately committed source/cohort binding with a passed
gate and exact input hashes. Functions can be tested with synthetic workbooks
without opening a real source workbook.
"""
from __future__ import annotations

import argparse
import collections
import gzip
import hashlib
import itertools
import json
import math
from pathlib import Path
import re
import shutil
import struct
import subprocess
import sys

import numpy as np

CELLS = ["K562", "HepG2"]
SCOPES = ["paired", *CELLS]
MODELS = ["zero", "I", "S"]
METRICS = ["mse", "mae"]
SEED = 0x2026102801
REPLICATES = 2000
MAX_INDICES = 50_000_000
TOLERANCE = 1e-10
PROTOCOL_SHA = "811930c2d3a5ee763014ea611acda4d2c84404b7d5ae399c7843c62d0a828149"
MODEL_SHA = "2ea476bc5413dc13486dec9f3d7e9162a50a0e47f1d83ce29efabc5f5cc5af9d"
DNA = "ACGT"
TRANS = str.maketrans("ACGT", "TGCA")
VARIANT_ID = re.compile(r"^(chr(?:[1-9]|1[0-9]|2[0-2]|X|Y)):([1-9][0-9]*):([ACGT]):([ACGT])$")


def require(condition, message):
    if not condition:
        raise ValueError(message)


def sha256(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def float_bits(value):
    return struct.pack("<d", float(value)).hex()


def rc(sequence):
    return sequence.translate(TRANS)[::-1]


TRIMERS = sorted({min("".join(t), rc("".join(t))) for t in itertools.product(DNA, repeat=3)})
FEATURES = ["rc3mer_" + t for t in TRIMERS] + ["gc_fraction"]


def exact_counts(sequence):
    require(isinstance(sequence, str) and len(sequence) == 200 and set(sequence) <= set(DNA),
            "expected exact 200nt uppercase ACGT core")
    histogram = collections.Counter(min(sequence[i:i + 3], rc(sequence[i:i + 3]))
                                    for i in range(198))
    return np.asarray([histogram[word] for word in TRIMERS]
                      + [sequence.count("C") + sequence.count("G")], dtype=np.int64)


def finite_source_number(cell, missing_tokens):
    """Only explicitly bound missing strings/blanks may exclude a required ID."""
    if cell.value is None:
        require(None in missing_tokens, "undocumented blank source cell")
        return None
    if cell.data_type == "s" and isinstance(cell.value, str) and cell.value in missing_tokens:
        return None
    require(cell.data_type == "n" and type(cell.value) in (int, float)
            and math.isfinite(cell.value), "source effect is not an ordinary finite numeric cell")
    return float(cell.value)


def verify_source_file(path, commitment):
    path = Path(path)
    require(path.is_file() and not path.is_symlink(), "missing/symlinked original workbook")
    require(path.stat().st_size == commitment["bytes"] and sha256(path) == commitment["sha256"],
            "original workbook byte commitment differs")


def checked_sheet(path, commitment, sheet_name, header_row, expected_columns):
    """Openpyxl keeps formulas visible; never trust cached formula values."""
    import openpyxl
    verify_source_file(path, commitment)
    book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        require(sheet_name in book.sheetnames, "bound source worksheet absent")
        sheet = book[sheet_name]
        # Physical rows, not a producer-supplied dimension estimate, define data.
        sheet.reset_dimensions()
        header = next(sheet.iter_rows(min_row=header_row, max_row=header_row,
                                     max_col=max(expected_columns)))
        for column, label in expected_columns.items():
            cell = header[column - 1]
            require(cell.data_type in {"s", "inlineStr"} and cell.value == label,
                    "bound textual worksheet header differs")
        return book, sheet
    except BaseException:
        book.close()
        raise


def read_s2_dna(path, commitment, required_ids):
    required_ids = set(required_ids)
    require(required_ids and all(isinstance(v, str) and v for v in required_ids), "empty/invalid cohort")
    book, sheet = checked_sheet(path, commitment, "S2.Variant Oligo Seqs", 2,
        {1: "Variant", 2: "Allele 1 Oligo", 3: "Allele 2 Oligo"})
    found, seen = {}, set()
    try:
        for number, row in enumerate(sheet.iter_rows(min_row=3, max_col=3), 3):
            identifier = row[0].value
            if identifier is None and all(cell.value is None for cell in row):
                continue
            require(row[0].data_type in {"s", "inlineStr"} and isinstance(identifier, str)
                    and identifier == identifier.strip() and identifier and identifier not in seen,
                    "invalid or duplicate exact source design ID")
            seen.add(identifier)
            if identifier not in required_ids:
                continue
            require(all(cell.data_type in {"s", "inlineStr"} for cell in row[1:]),
                    "DNA must be literal strings, never formulas")
            reference, alternate = row[1].value, row[2].value
            before, after = exact_counts(reference), exact_counts(alternate)
            different = [i for i in range(200) if reference[i] != alternate[i]]
            match = VARIANT_ID.fullmatch(identifier)
            require(match is not None and different == [99] and match[3] == reference[99]
                    and match[4] == alternate[99], "source ID/alleles/central SNV contradict")
            found[identifier] = {"reference": reference, "alternate": alternate,
                "source_row_1based": number, "reference_counts": before.tolist(),
                "alternate_counts": after.tolist(), "delta_counts": (after - before).tolist(),
                "gc_preserving": ((reference[99] in "GC") == (alternate[99] in "GC"))}
    finally:
        book.close()
    require(set(found) == required_ids, "required metadata cohort missing from original S2")
    return found, {"source_unique_ids": len(seen), "required_ids": len(found)}


def read_s4_effects(path, commitment, required_ids, missing_tokens, absence_allowed=False,
                    cell_ledger=None):
    required_ids = set(required_ids)
    require(required_ids, "no required source IDs")
    book, sheet = checked_sheet(path, commitment, "S4.MPRA results", 2,
        {1: "Variant", 40: "log2Skew in HEPG2", 41: "log2Skew in K562"})
    found, missing, seen = {}, {}, set()
    try:
        for number, row in enumerate(sheet.iter_rows(min_row=3, max_col=41), 3):
            identifier = row[0].value
            if identifier is None and all(row[i].value is None for i in (0, 39, 40)):
                continue
            require(row[0].data_type in {"s", "inlineStr"} and isinstance(identifier, str)
                    and identifier and identifier == identifier.strip() and identifier not in seen,
                    "invalid or duplicate exact source outcome ID")
            seen.add(identifier)
            if identifier not in required_ids:
                continue
            observed = {}
            for cell, column in [("K562", 41), ("HepG2", 40)]:
                source_cell = row[column - 1]
                audit = {"id": identifier, "context": cell, "source_row_1based": number,
                    "cell_data_type": source_cell.data_type,
                    "python_type": type(source_cell.value).__name__,
                    "openpyxl_value_repr_not_raw_XML_lexeme": repr(source_cell.value)}
                if cell_ledger is not None:
                    cell_ledger.append(audit)
                try:
                    observed[cell] = finite_source_number(source_cell, missing_tokens)
                    audit.update(status="ordinary_numeric" if observed[cell] is not None else "bound_missing",
                        float64_le_hex=None if observed[cell] is None else float_bits(observed[cell]))
                except ValueError as error:
                    audit.update(status="gate_failure", reason=str(error))
                    raise
            absent_cells = [cell for cell, value in observed.items() if value is None]
            record = {"source_row_1based": number, "observed": observed,
                "observed_float64_little_endian": {cell: None if value is None else float_bits(value)
                                                   for cell, value in observed.items()}}
            if absent_cells:
                missing[identifier] = record | {"missing_contexts": absent_cells}
            else:
                found[identifier] = record
    finally:
        book.close()
    absent = sorted(required_ids - set(found) - set(missing))
    require(absence_allowed or not absent, "required source IDs absent without frozen source-QC rule")
    return found, {"incomplete": missing, "absent_ids": absent,
        "all_source_unique_ids": len(seen), "complete_ids": len(found)}


def read_saved_fits(core, expected_manifest_sha=MODEL_SHA):
    core = Path(core)
    require(sha256(core / "artifacts.json") == expected_manifest_sha, "frozen fit manifest differs")
    manifest = json.loads((core / "artifacts.json").read_bytes())
    require(manifest["schema_version"] == 5, "wrong frozen artifact schema")
    for name, expected in manifest["files"].items():
        member = core / name
        require(not member.is_symlink() and member.resolve().is_relative_to(core.resolve())
                and sha256(member) == expected, "frozen fit artifact commitment differs")
    with gzip.open(core / "study.json.gz", "rt") as handle:
        study = json.load(handle)
    require(len(study["folds"]) == 5, "need all five frozen fits")
    result = {}
    for model in ("I", "S"):
        result[model] = {}
        for cell in CELLS:
            slopes = []
            for fold in study["folds"]:
                fitted = fold["fit"]
                require(fitted["feature_names"] == FEATURES, "frozen feature order differs")
                ridge = fitted["shared_dna_fit"] if model == "S" else fitted["independent_context_fits"][cell]
                beta = np.asarray(ridge["standardized_coefficients"], dtype=np.float64)
                scale = np.asarray(ridge["training_feature_scales"], dtype=np.float64)
                require(beta.shape == scale.shape == (33,) and np.isfinite(beta).all()
                        and np.isfinite(scale).all() and (scale > 0).all(), "invalid saved slopes")
                slopes.append(beta / scale)
            result[model][cell] = np.stack(slopes)
    return result


def predict_fixed(dna_rows, slopes):
    identifiers = sorted(dna_rows)
    counts = np.asarray([dna_rows[i]["delta_counts"] for i in identifiers], dtype=np.int64)
    delta = counts / np.asarray([198] * 32 + [200], dtype=np.float64)
    predictions, fold_predictions = {}, {}
    for model in ("I", "S"):
        predictions[model], fold_predictions[model] = {}, {}
        for cell in CELLS:
            # Independent NumPy matrix products, rather than primary fsum loops.
            per_fold = delta @ slopes[model][cell].T
            fold_predictions[model][cell] = per_fold
            predictions[model][cell] = per_fold.mean(axis=1)
    return identifiers, predictions, fold_predictions


def splitmix_indices(groups, replicates=REPLICATES, seed=SEED):
    require(type(groups) is int and groups > 0 and 0 < replicates <= 100000
            and groups * replicates <= MAX_INDICES, "invalid resampling resource contract")
    state = seed
    mask = (1 << 64) - 1
    result = np.empty((replicates, groups), dtype=np.int64)
    for r in range(replicates):
        for j in range(groups):
            state = (state + 11400714819323198485) & mask
            value = state
            value = ((value ^ (value >> 30)) * 13787848793156543929) & mask
            value = ((value ^ (value >> 27)) * 10723151780598845931) & mask
            result[r, j] = (value ^ (value >> 31)) % groups
    return result


def type7(values, probabilities=(0.025, 0.975)):
    return np.quantile(np.asarray(values, dtype=np.float64), probabilities, method="linear").tolist()


def loss_object(values):
    # Tensor axes are scope, model, metric.
    return {scope: {model: {metric: float(values[s, m, q]) for q, metric in enumerate(METRICS)}
                    for m, model in enumerate(MODELS)} for s, scope in enumerate(SCOPES)}


def gain_tensor(values):
    return values[:, :1, :] - values[:, 1:, :]


def gain_object(values):
    return {scope: {model: {metric: float(values[s, m, q]) for q, metric in enumerate(METRICS)}
                    for m, model in enumerate(("I", "S"))} for s, scope in enumerate(SCOPES)}


def association(observed, predicted):
    observed, predicted = np.asarray(observed), np.asarray(predicted)
    signs = ("negative", "zero", "positive")
    a = np.sign(observed).astype(np.int64) + 1
    p = np.sign(predicted).astype(np.int64) + 1
    confusion = np.zeros((3, 3), dtype=np.int64)
    np.add.at(confusion, (a, p), 1)
    if np.all(observed == observed[0]) or np.all(predicted == predicted[0]):
        correlation = {"estimable": False, "value": None, "reason": "constant observed or predicted vector"}
    else:
        correlation = {"estimable": True, "value": float(np.corrcoef(observed, predicted)[0, 1]), "reason": None}
    matching = int(np.trace(confusion))
    return {"observations": observed.size, "pearson": correlation,
        "exact_sign_agreement_including_zero": matching / observed.size,
        "sign_agreement_count": matching,
        "sign_confusion_observed_then_predicted": {x: {y: int(confusion[i, j]) for j, y in enumerate(signs)}
                                                  for i, x in enumerate(signs)},
        "exact_zero_observations": int(np.count_nonzero(observed == 0)),
        "exact_zero_predictions": int(np.count_nonzero(predicted == 0))}


def evaluate(variants, replicates=REPLICATES, seed=SEED):
    require(variants and [r["id"] for r in variants] == sorted({r["id"] for r in variants}),
            "variant order must be unique and lexical")
    targets = np.asarray([[r["observed"][c] for c in CELLS] for r in variants], dtype=np.float64)
    forecasts = np.asarray([[[0.0, r["predictions"]["I"][c], r["predictions"]["S"][c]]
                            for c in CELLS] for r in variants], dtype=np.float64)
    require(np.isfinite(targets).all() and np.isfinite(forecasts).all(), "nonfinite input")
    error = forecasts - targets[:, :, None]
    base_losses = np.stack((error ** 2, np.abs(error)), axis=-1)
    require(np.isfinite(base_losses).all(), "loss overflow")
    # Retain separate paired/cell scopes, averaging cells only for paired.
    loss = np.concatenate((base_losses.mean(axis=1, keepdims=True), base_losses), axis=1)
    member_map, locus_map = collections.defaultdict(list), collections.defaultdict(list)
    for i, row in enumerate(variants):
        member_map[row["component"]].append(i)
        locus_map[row["locus"]].append(i)
    groups = sorted(member_map)
    members = [np.asarray(member_map[group], dtype=np.int64) for group in groups]
    component_losses = np.stack([loss[index].mean(axis=0) for index in members])
    sizes = np.asarray([len(index) for index in members], dtype=np.int64)
    observed = {}
    for weighting, values in [("equal_component", component_losses.mean(axis=0)),
                              ("equal_variant", np.average(component_losses, weights=sizes, axis=0))]:
        observed[weighting] = {"losses": loss_object(values), "gains_vs_zero": gain_object(gain_tensor(values))}
    ledger = [{"id": row["id"], "component": row["component"], "locus": row["locus"],
        "gc_preserving": row["gc_preserving"], "losses": loss_object(loss[i]),
        "observed_float64_little_endian": {cell: float_bits(row["observed"][cell]) for cell in CELLS},
        "prediction_float64_little_endian": {model: {cell: float_bits(row["predictions"][model][cell])
            for cell in CELLS} for model in ("I", "S")}} for i, row in enumerate(variants)]
    components = [{"component": name, "variant_count": len(indices),
        "variant_ids": [variants[i]["id"] for i in indices],
        "loci": sorted({variants[i]["locus"] for i in indices}), "losses": loss_object(component_losses[g])}
        for g, (name, indices) in enumerate(zip(groups, members))]
    draw_indices = splitmix_indices(len(groups), replicates, seed)
    statistics = []
    sampled = {w: {"losses": [], "gains_vs_zero": []} for w in ("equal_component", "equal_variant")}
    for indices in draw_indices:
        record = {}
        for weighting in sampled:
            selected = component_losses[indices]
            values = (selected.mean(axis=0) if weighting == "equal_component" else
                      np.average(selected, weights=sizes[indices], axis=0))
            gains = gain_tensor(values)
            sampled[weighting]["losses"].append(values)
            sampled[weighting]["gains_vs_zero"].append(gains)
            record[weighting] = {"losses": loss_object(values), "gains_vs_zero": gain_object(gains)}
        statistics.append(record)
    intervals = {}
    for weighting in sampled:
        intervals[weighting] = {}
        for kind, models in [("losses", MODELS), ("gains_vs_zero", ["I", "S"])]:
            values = np.asarray(sampled[weighting][kind])
            intervals[weighting][kind] = {scope: {model: {metric: type7(values[:, s, m, q])
                for q, metric in enumerate(METRICS)} for m, model in enumerate(models)}
                for s, scope in enumerate(SCOPES)}
    gc_diagnostics = {}
    for boolean, label in [(True, "GC_preserving"), (False, "GC_changing")]:
        indices = [i for i, row in enumerate(variants) if row["gc_preserving"] is boolean]
        represented = sorted({variants[i]["component"] for i in indices})
        selected_set = set(indices)
        means = [loss[[i for i in member_map[group] if i in selected_set]].mean(axis=0)
                 for group in represented]
        gc_diagnostics[label] = {"variant_ids": [variants[i]["id"] for i in indices],
            "represented_component_ids": represented,
            "empty_original_component_ids": sorted(set(groups) - set(represented)),
            "variants": len(indices), "estimable": bool(indices),
            "equal_variant_losses": loss_object(loss[indices].mean(axis=0)) if indices else None,
            "equal_represented_component_losses": loss_object(np.mean(means, axis=0)) if indices else None}
    loci = []
    for name in sorted(locus_map):
        indices = locus_map[name]
        loci.append({"locus": name, "variant_ids": [variants[i]["id"] for i in indices],
            "component_ids": sorted({variants[i]["component"] for i in indices}),
            "variant_count": len(indices), "equal_variant_losses": loss_object(loss[indices].mean(axis=0))})
    associations = {}
    for scope in SCOPES:
        selected_cells = [0, 1] if scope == "paired" else [CELLS.index(scope)]
        associations[scope] = {model: association(targets[:, selected_cells].ravel(),
            forecasts[:, selected_cells, m].ravel()) for m, model in enumerate(MODELS)}
    gain = observed["equal_component"]["gains_vs_zero"]["paired"]["I"]["mse"]
    interval = intervals["equal_component"]["gains_vs_zero"]["paired"]["I"]["mse"]
    support = len(variants) >= 100 and len(groups) >= 10
    return {"schema_version": 1, "variant_order": [row["id"] for row in variants],
        "contexts": CELLS, "model_order": MODELS, "variant_count": len(variants),
        "component_count": len(groups), "observed": observed,
        "primary": {"metric": "equal_component_macroMSE_zero_minus_I", "observed_improvement": gain,
            "conditional_95_percent_range": interval, "minimum_support_passed": support,
            "positive_observed_gain": gain > 0, "positive_lower_endpoint": interval[0] > 0,
            "useful_transfer_rule_passed": support and gain > 0 and interval[0] > 0},
        "associations_equal_observation_weight": associations,
        "gc_diagnostics_descriptive_only": gc_diagnostics, "loci": loci,
        "per_variant": ledger, "per_component": components,
        "resampling": {"component_order": groups, "replicates": replicates, "seed": seed,
            "rng": "SplitMix64", "index_rule": "next_u64 modulo component_count",
            "interval_probabilities": [0.025, 0.975], "quantile_method": "Type7",
            "resampling_indices": draw_indices.tolist(), "sampled_statistics": statistics,
            "intervals": intervals},
        "limits": ["Independent NumPy arithmetic; no new fit or outcome calibration.",
            "An empirical run additionally requires raw-workbook source/eligibility checks and the second committed binding.",
            "Metadata components must be separately authenticated; this evaluator does not rediscover genome homology."]}


def compare_trees(primary, independent, tolerance=TOLERANCE):
    checks, failures = [], []
    maximum_error, numeric_cells = 0.0, 0

    def walk(a, b, path):
        nonlocal maximum_error, numeric_cells
        if path == "/limits":
            return
        if "prediction_float64_little_endian" in path and isinstance(a, str) and isinstance(b, str):
            a, b = struct.unpack("<d", bytes.fromhex(a))[0], struct.unpack("<d", bytes.fromhex(b))[0]
        if type(a) is float and type(b) is float:
            numeric_cells += 1
            error = abs(a - b)
            maximum_error = max(maximum_error, error)
            if not math.isfinite(error) or error > tolerance:
                failures.append({"path": path, "primary": a, "independent": b, "error": error})
        elif isinstance(a, dict) and isinstance(b, dict):
            if set(a) != set(b):
                failures.append({"path": path, "reason": "object fields differ"})
            for key in sorted(set(a) & set(b)):
                walk(a[key], b[key], path + "/" + key)
        elif isinstance(a, list) and isinstance(b, list):
            if len(a) != len(b):
                failures.append({"path": path, "reason": "array length differs"})
            for i, (x, y) in enumerate(zip(a, b)):
                walk(x, y, path + "/" + str(i))
        elif type(a) != type(b) or a != b:
            failures.append({"path": path, "primary": a, "independent": b})

    for key in sorted(set(primary) | set(independent)):
        before = len(failures)
        if key not in primary or key not in independent:
            failures.append({"path": "/" + key, "reason": "missing top-level field"})
        else:
            walk(primary[key], independent[key], "/" + key)
        checks.append({"name": key, "matched": before == len(failures)})
    return {"matched": not failures, "checks": checks, "failure_count": len(failures),
        "failures": failures, "max_absolute_numeric_error": maximum_error,
        "numeric_cells_compared": numeric_cells, "absolute_tolerance": tolerance,
        "observed_float_bits_and_integer_indices_exact": not any(
            "observed_float64_little_endian" in f["path"] or "resampling_indices" in f["path"]
            for f in failures)}


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        require(key not in result, "duplicate JSON key")
        result[key] = value
    return result


def read_json(path):
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8") as handle:
        return json.load(handle, object_pairs_hook=unique_object,
            parse_constant=lambda token: (_ for _ in ()).throw(ValueError("nonfinite JSON literal")))


def write_json(path, value):
    data = (json.dumps(value, sort_keys=True, indent=2, allow_nan=False) + "\n").encode()
    if str(path).endswith(".gz"):
        with Path(path).open("wb") as raw:
            with gzip.GzipFile(fileobj=raw, filename="", mode="wb", mtime=0) as handle:
                handle.write(data)
    else:
        Path(path).write_bytes(data)


def manifest_directory(root):
    files = {str(path.relative_to(root)): {"sha256": sha256(path), "bytes": path.stat().st_size}
             for path in sorted(root.rglob("*")) if path.is_file() and path.name != "manifest.json"}
    write_json(root / "manifest.json", {"schema_version": 1, "files": files})


def check_committed_input(repo_root, commit, path, expected_sha):
    require(isinstance(commit, str) and re.fullmatch(r"[0-9a-f]{40}", commit) is not None,
            "a full second-freeze commit SHA is required")
    repo_root, path = Path(repo_root).resolve(), Path(path).resolve()
    require(path.is_relative_to(repo_root), "second-freeze input must be in the checkout")
    member = path.relative_to(repo_root).as_posix()
    ancestor = subprocess.run(["git", "-C", str(repo_root), "merge-base", "--is-ancestor", commit, "HEAD"],
                              stdout=subprocess.DEVNULL, stderr=subprocess.PIPE)
    require(ancestor.returncode == 0, "second-freeze commit is not an ancestor of current HEAD")
    with subprocess.Popen(["git", "-C", str(repo_root), "show", commit + ":" + member],
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE) as process:
        digest = hashlib.file_digest(process.stdout, "sha256").hexdigest()
        process.stdout.close()
        error = process.stderr.read()
        result = process.wait()
    require(result == 0 and digest == expected_sha, "input bytes were not frozen in the named commit: " + member)


def validate_empirical_binding(args):
    # This gate runs before importing openpyxl or opening any source ZIP.
    require(sha256(args.binding) == args.expected_binding_sha256, "binding SHA256 differs")
    binding = read_json(args.binding)
    require(binding.get("schema_version") == 1 and binding.get("protocol_id") == "variant-transfer-v1"
            and binding.get("protocol_sha256") == PROTOCOL_SHA, "wrong semantic protocol")
    require(binding.get("source_cohort_gate", {}).get("passed") is True,
            "sealed source/cohort gate is not passed; no raw outcome workbook will be opened")
    require(binding.get("model_artifacts_sha256") == MODEL_SHA, "wrong frozen activity fits")
    require(sha256(args.protocol) == PROTOCOL_SHA, "semantic protocol changed")
    sheet = binding.get("sheet", {})
    require(sheet.get("name") == "S4.MPRA results" and sheet.get("header_row_1based") == 2
        and sheet.get("headers") == {"A": "Variant", "AN": "log2Skew in HEPG2", "AO": "log2Skew in K562"},
        "unexpected exact source sheet/header binding")
    require(binding.get("effects") == {"HepG2": "AN", "K562": "AO"}, "effect column orientation differs")
    tokens = binding.get("missing_policy", {}).get("allowed_tokens")
    require(tokens in ([], ["NA"]), "missing tokens are not a frozen exact source policy")
    require(not tokens or binding["missing_policy"].get("evidence"), "NA requires source documentation")
    require(binding.get("cohort", {}).get("format") in ("json", "jsonl.gz")
            and type(binding["cohort"].get("variant_count")) is int
            and binding["cohort"]["variant_count"] > 0, "invalid frozen metadata cohort contract")
    require(sha256(args.cohort) == binding["cohort"]["sha256"], "metadata cohort bytes differ")
    for path, pin in ((args.binding, args.expected_binding_sha256), (args.protocol, PROTOCOL_SHA),
                      (args.cohort, binding["cohort"]["sha256"])):
        check_committed_input(args.repo_root, args.freeze_commit, path, pin)
    protocol = read_json(args.protocol)
    require(protocol["contexts"]["primary"] == CELLS, "frozen cell order differs")
    require(protocol["analysis_contract"] == {"contexts": CELLS,
        "primary_weighting": "equal_component", "baseline": "zero", "candidate": "I",
        "bootstrap": {"replicates": REPLICATES, "rng": "SplitMix64", "seed": SEED,
            "quantiles": [0.025, 0.975], "quantile_type": 7},
        "minimum_variants": 100, "minimum_components": 10}, "fixed analysis contract differs")
    verify_source_file(args.s4_workbook, binding["source"])
    verify_source_file(args.s2_workbook, protocol["source"]["design_table"])
    return protocol, binding


def load_frozen_cohort(path, binding):
    if binding["cohort"]["format"] == "jsonl.gz":
        with gzip.open(path, "rt", encoding="utf-8") as handle:
            rows = [json.loads(line, object_pairs_hook=unique_object) for line in handle]
    else:
        payload = read_json(path)
        require(set(payload) == {"schema_version", "pairs"} and payload["schema_version"] == 1,
                "invalid frozen cohort object")
        rows = payload["pairs"]
    require(len(rows) == binding["cohort"]["variant_count"], "frozen metadata cohort count differs")
    lookup = {}
    for row in rows:
        require(all(isinstance(row.get(k), str) and row[k] and row[k] == row[k].strip()
                    for k in ("id", "component", "locus")) and row["id"] not in lookup,
                "duplicate or invalid frozen cohort labels")
        before, after = exact_counts(row["sequence_ref"]), exact_counts(row["sequence_alt"])
        changed = [i for i in range(200) if row["sequence_ref"][i] != row["sequence_alt"][i]]
        preserving = (row["sequence_ref"][99] in "GC") == (row["sequence_alt"][99] in "GC")
        require(changed == [99] and type(row.get("position_0based", 99)) is int
                and row.get("position_0based", 99) == 99 and type(row.get("gc_preserving")) is bool
                and row["gc_preserving"] == preserving, "cohort SNV/GC metadata contradict")
        lookup[row["id"]] = row
    return lookup


def run_empirical(args):
    require(not args.out.exists(), "preserve all previous empirical attempts")
    protocol, binding = validate_empirical_binding(args)
    cohort = load_frozen_cohort(args.cohort, binding)
    args.out.mkdir(parents=True)
    for name, path in (("implementation.py", Path(__file__)), ("binding.json", args.binding),
                       ("protocol.json", args.protocol)):
        shutil.copyfile(path, args.out / name)
    write_json(args.out / "metadata-cohort.json.gz", cohort)
    write_json(args.out / "inputs.json", {"protocol_sha256": PROTOCOL_SHA,
        "binding_sha256": args.expected_binding_sha256, "second_freeze_commit": args.freeze_commit,
        "cohort_sha256": binding["cohort"]["sha256"], "S4": binding["source"],
        "S2": protocol["source"]["design_table"], "frozen_model_artifacts_sha256": MODEL_SHA,
        "numpy_version": np.__version__, "implementation_sha256": sha256(Path(__file__)),
        "component_partition_authenticated_not_independently_reconstructed": True,
        "models_refitted": False})
    cells = []
    try:
        dna, dna_report = read_s2_dna(args.s2_workbook, protocol["source"]["design_table"], cohort)
        for identifier, row in dna.items():
            require(row["reference"] == cohort[identifier]["sequence_ref"]
                    and row["alternate"] == cohort[identifier]["sequence_alt"]
                    and row["gc_preserving"] == cohort[identifier]["gc_preserving"],
                    "original S2 DNA contradicts frozen metadata")
        write_json(args.out / "source-DNA.json.gz", dna)
        ids, predictions, fold_predictions = predict_fixed(dna, read_saved_fits(args.model_artifacts))
        write_json(args.out / "fixed-predictions.json.gz", {"variant_order": ids,
            "predictions": {m: {c: values.tolist() for c, values in group.items()}
                            for m, group in predictions.items()},
            "fold_predictions": {m: {c: values.tolist() for c, values in group.items()}
                                 for m, group in fold_predictions.items()}})
        # Numeric S4 cells are first accessed here, after every freeze/model/DNA check.
        observed, missing = read_s4_effects(args.s4_workbook, binding["source"], cohort,
            binding["missing_policy"]["allowed_tokens"], cell_ledger=cells)
        write_json(args.out / "source-observations.json.gz", observed)
        write_json(args.out / "source-missingness.json.gz", missing)
        index = {identifier: i for i, identifier in enumerate(ids)}
        variants = [{"id": identifier, "component": cohort[identifier]["component"],
            "locus": cohort[identifier]["locus"], "gc_preserving": dna[identifier]["gc_preserving"],
            "observed": observed[identifier]["observed"],
            "predictions": {model: {cell: float(predictions[model][cell][index[identifier]])
                                      for cell in CELLS} for model in ("I", "S")}}
                    for identifier in sorted(observed)]
        require(variants, "no complete variants; retain source missingness without an estimable comparison")
        write_json(args.out / "independent-normalized.json.gz", {"schema_version": 1, "variants": variants})
        result = evaluate(variants)
        write_json(args.out / "evaluation.json.gz", result)
        write_json(args.out / "summary.json", {key: value for key, value in result.items()
            if key not in ("per_variant", "resampling")})
        write_json(args.out / "status.json", {"status": "complete", "metadata_variants": len(cohort),
            "complete_case_variants": len(variants), "incomplete_variants": len(missing["incomplete"]),
            "source_DNA_inventory": dna_report, "source_outcome_unique_IDs": missing["all_source_unique_ids"],
            "protocol_sha256": PROTOCOL_SHA, "binding_sha256": args.expected_binding_sha256,
            "source_float_bits_preserved": True,
            "component_partition_authenticated_not_independently_reconstructed": True,
            "full_protocol_independent_reproduction_claimed": False})
        print(json.dumps({"complete_variants": len(variants), "primary": result["primary"]}))
    except BaseException as error:
        write_json(args.out / "failure.json", {"error_type": type(error).__name__, "error": str(error),
            "protocol_sha256": PROTOCOL_SHA, "binding_sha256": args.expected_binding_sha256})
        raise
    finally:
        write_json(args.out / "source-cell-type-ledger.json.gz", cells)
        manifest_directory(args.out)


def verify_run_directory(root):
    root = Path(root)
    require(root.is_dir() and not root.is_symlink(), "missing/symlinked run directory")
    manifest = read_json(root / "manifest.json")
    expected = manifest["files"]
    actual = {str(p.relative_to(root)): p for p in root.rglob("*") if p.is_file()
              and p.relative_to(root).as_posix() != "manifest.json"}
    require(set(actual) == set(expected), "run manifest inventory differs")
    for name, commitment in expected.items():
        path = actual[name]
        require(not path.is_symlink() and path.resolve().is_relative_to(root.resolve()), "unsafe run file")
        require(sha256(path) == (commitment if isinstance(commitment, str) else commitment["sha256"]),
                "run file SHA differs")
        if isinstance(commitment, dict):
            require(path.stat().st_size == commitment["bytes"], "run file bytes differ")
    return {"sha256": sha256(root / "manifest.json"), "bytes": (root / "manifest.json").stat().st_size}


def compare_raw_sources_and_predictions(own_root, observations_root, prediction_roots):
    own_root, observations_root = Path(own_root), Path(observations_root)
    parents = {"independent": verify_run_directory(own_root),
               "primary_observations": verify_run_directory(observations_root)}
    dna = read_json(own_root / "source-DNA.json.gz")
    metadata = read_json(own_root / "metadata-cohort.json.gz")
    source = read_json(own_root / "source-observations.json.gz")
    missing = read_json(own_root / "source-missingness.json.gz")
    fixed = read_json(own_root / "fixed-predictions.json.gz")
    primary_observations = read_json(observations_root / "observations.json.gz")
    inputs = read_json(own_root / "inputs.json")
    require(primary_observations["protocol_sha256"] == inputs["protocol_sha256"] == PROTOCOL_SHA
            and primary_observations["binding_sha256"] == inputs["binding_sha256"],
            "source comparison does not identify the same committed contract")
    expected_ids = set(metadata)
    require(set(dna) == expected_ids and set(source) | set(missing["incomplete"]) == expected_ids
            and not missing["absent_ids"], "independent metadata/source coverage differs")
    rows = primary_observations["variants"]
    require(len(rows) == len({r["id"] for r in rows}) and {r["id"] for r in rows} == expected_ids,
            "primary observation cohort differs")
    checks, failures = [], []
    source_bits, exact_numerators, prediction_cells, maximum_error = 0, 0, 0, 0.0

    def check(name, condition, detail=None):
        checks.append({"name": name, "matched": bool(condition)})
        if not condition:
            failures.append({"name": name, "detail": detail})

    def number(name, primary, independent):
        nonlocal prediction_cells, maximum_error
        require(type(primary) in (int, float) and type(independent) in (int, float), "nonnumeric prediction")
        error = abs(primary - independent)
        prediction_cells += 1
        maximum_error = max(maximum_error, error)
        if not math.isfinite(error) or error > TOLERANCE:
            failures.append({"name": name, "primary": primary, "independent": independent, "error": error})

    # Each failure records its exact source identity; counts are not hidden by a summary.
    before = len(failures)
    for row in rows:
        identifier = row["id"]
        actual = source.get(identifier, missing["incomplete"].get(identifier))
        for cell in CELLS:
            source_bits += 1
            if row["source_float64_le_hex"][cell] != actual["observed_float64_little_endian"][cell]:
                failures.append({"name": "source_float_bits", "id": identifier, "context": cell})
        expected_complete = identifier in source
        expected_missing = [] if expected_complete else missing["incomplete"][identifier]["missing_contexts"]
        if row["complete_case"] is not expected_complete or row["missing_contexts"] != expected_missing:
            failures.append({"name": "source_complete_case_rule", "id": identifier})
        if any(row[key] != metadata[identifier][key] for key in ("component", "locus", "gc_preserving")):
            failures.append({"name": "source_frozen_metadata", "id": identifier})
        if (row["reference"] != dna[identifier]["reference"]
                or row["alternate"] != dna[identifier]["alternate"]):
            failures.append({"name": "original_S2_DNA", "id": identifier})
    checks.append({"name": "all_original_S2_DNA_metadata_source_float_bits_and_complete_cases", "matched": before == len(failures)})
    indices = {identifier: i for i, identifier in enumerate(fixed["variant_order"])}
    for model, directory in prediction_roots.items():
        parents["prediction_" + model] = verify_run_directory(directory)
        payload = read_json(Path(directory) / "predictions.json")
        check(model + "_frozen_model_feature_contract", payload["model"] == model
            and payload["feature_names"] == FEATURES
            and payload["feature_denominators"] == [198] * 32 + [200]
            and payload["commitments"]["artifacts_sha256"] == MODEL_SHA)
        expected_pairs = {identifier + "::" + cell for identifier in source for cell in CELLS}
        primary_pairs = payload["pairs"]
        require(len(primary_pairs) == len({r["id"] for r in primary_pairs})
                and {r["id"] for r in primary_pairs} == expected_pairs, "prediction pair cohort differs")
        before = len(failures)
        for row in primary_pairs:
            identifier, cell = row["id"].rsplit("::", 1)
            require(cell in CELLS and row["context"] == cell and row["position_0based"] == 99,
                    "prediction context/position differs")
            data = dna[identifier]
            for primary_key, own_key in (("reference_feature_numerators", "reference_counts"),
                ("alternate_feature_numerators", "alternate_counts"), ("feature_numerator_delta", "delta_counts")):
                values = row[primary_key]
                exact_numerators += 33
                if len(values) != 33 or any(type(v) is not int for v in values) or values != data[own_key]:
                    failures.append({"name": "exact_feature_numerators", "id": row["id"], "model": model,
                                     "field": primary_key})
            for field, seq in (("reference_sha256", data["reference"]), ("alternate_sha256", data["alternate"])):
                if row[field] != hashlib.sha256(seq.encode()).hexdigest():
                    failures.append({"name": "prediction_sequence_hash", "id": row["id"], "model": model})
            i = indices[identifier]
            number(model + "/" + row["id"] + "/ensemble", row["predicted_alternate_minus_reference"],
                   fixed["predictions"][model][cell][i])
            require(len(row["fold_predictions"]) == 5, "missing fixed fold prediction")
            for f, value in enumerate(row["fold_predictions"]):
                number(model + "/" + row["id"] + "/fold/" + str(f), value,
                       fixed["fold_predictions"][model][cell][i][f])
        checks.append({"name": model + "_all_integer_numerators_DNA_hashes_folds_and_ensembles", "matched": before == len(failures)})
    return {"matched": not failures, "checks": checks, "failures": failures, "failure_count": len(failures),
        "parent_manifests": parents, "source_float_bits_exact_comparisons": source_bits,
        "exact_integer_feature_numerators_compared": exact_numerators,
        "fixed_prediction_numbers_compared": prediction_cells, "maximum_absolute_prediction_error": maximum_error,
        "absolute_tolerance": TOLERANCE,
        "component_partition_authenticated_not_independently_reconstructed": True,
        "full_protocol_independent_reproduction_claimed": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--compare-primary", type=Path)
    parser.add_argument("--independent-evaluation", type=Path)
    parser.add_argument("--independent-run", type=Path)
    parser.add_argument("--primary-observations", type=Path)
    parser.add_argument("--prediction-I", type=Path)
    parser.add_argument("--prediction-S", type=Path)
    parser.add_argument("--run", action="store_true")
    parser.add_argument("--binding", type=Path)
    parser.add_argument("--expected-binding-sha256")
    parser.add_argument("--cohort", type=Path)
    parser.add_argument("--freeze-commit")
    parser.add_argument("--repo-root", type=Path, default=Path.cwd())
    parser.add_argument("--protocol", type=Path, default=Path("experiments/variant-transfer-v1.json"))
    parser.add_argument("--s2-workbook", type=Path)
    parser.add_argument("--s4-workbook", type=Path)
    parser.add_argument("--model-artifacts", type=Path, default=Path("research/results/context-activity-v3/core"))
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    require(not args.out.exists(), "preserve previous attempts")
    if args.run:
        require(all(getattr(args, name) is not None for name in (
            "binding", "expected_binding_sha256", "cohort", "freeze_commit", "s2_workbook", "s4_workbook")),
            "run requires the committed binding/cohort pins and original workbook paths")
        require(args.compare_primary is None and args.independent_evaluation is None,
                "empirical run and saved-result comparison are separate attempts")
        run_empirical(args)
        return 0
    if args.independent_run is not None:
        require(args.compare_primary is not None and args.compare_primary.is_dir()
                and args.primary_observations is not None and args.prediction_I is not None
                and args.prediction_S is not None and args.independent_evaluation is None,
                "full comparison requires primary evaluation/observations/predictions and independent run directories")
        args.out.mkdir(parents=True)
        shutil.copyfile(Path(__file__), args.out / "implementation.py")
        try:
            own_pin = verify_run_directory(args.independent_run)
            primary_pin = verify_run_directory(args.compare_primary)
            primary = read_json(args.compare_primary / "evaluation.json.gz")
            own = read_json(args.independent_run / "evaluation.json.gz")
            result = compare_trees(primary, own)
            normalized = compare_trees(
                {"variants": read_json(args.compare_primary / "normalized-input.json")["variants"]},
                {"variants": read_json(args.independent_run / "independent-normalized.json.gz")["variants"]})
            source = compare_raw_sources_and_predictions(args.independent_run, args.primary_observations,
                {"I": args.prediction_I, "S": args.prediction_S})
            result.update(primary_manifest=primary_pin, independent_manifest=own_pin,
                normalized_handoff_comparison=normalized, raw_source_and_fixed_prediction_comparison=source,
                component_partition_authenticated_not_independently_reconstructed=True,
                full_protocol_independent_reproduction_claimed=False)
            result["matched"] = result["matched"] and normalized["matched"] and source["matched"]
            result["failure_count"] += normalized["failure_count"] + source["failure_count"]
            result["checks"] += [{"name": "normalized_handoff", "matched": normalized["matched"]},
                                  {"name": "original_source_features_and_fixed_predictions", "matched": source["matched"]}]
            write_json(args.out / "comparison.json", result)
            print(json.dumps({"matched": result["matched"], "failure_count": result["failure_count"]}))
            return 0 if result["matched"] else 2
        except BaseException as error:
            write_json(args.out / "failure.json", {"error": str(error), "error_type": type(error).__name__})
            raise
        finally:
            manifest_directory(args.out)
    require(args.compare_primary is not None and args.independent_evaluation is not None,
            "Empirical reader remains gated pending second committed source/cohort binding; only saved-result comparison is enabled")
    with gzip.open(args.compare_primary, "rt") as stream:
        primary = json.load(stream)
    with gzip.open(args.independent_evaluation, "rt") as stream:
        own = json.load(stream)
    result = compare_trees(primary, own)
    args.out.mkdir(parents=True)
    (args.out / "comparison.json").write_text(json.dumps(result, indent=2, allow_nan=False) + "\n")
    print(json.dumps({"matched": result["matched"], "failures": result["failure_count"]}))
    return 0 if result["matched"] else 2


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except (ValueError, KeyError, TypeError, OSError, json.JSONDecodeError) as error:
        print("Independent variant transfer failed: " + str(error), file=sys.stderr)
        raise SystemExit(2)

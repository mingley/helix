"""Independent transfer controls use synthetic numbers and generated workbooks."""
import copy
import hashlib
import importlib.util
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import types
import unittest

import numpy as np
import openpyxl

SPEC = importlib.util.spec_from_file_location("independent_variant", Path(__file__).resolve().parents[2]
                                           / "scripts/replicate_variant_transfer.py")
M = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(M)


def rows():
    result = []
    for i, (component, observed, predicted) in enumerate([
        ("a", [1.0, 3.0], [0.5, 2.0]),
        ("b", [2.0, 4.0], [1.0, 2.0]),
        ("b", [0.0, -2.0], [0.0, -1.0]),
    ]):
        result.append({"id": str(i), "component": component, "locus": str(i),
            "gc_preserving": i == 0, "observed": dict(zip(M.CELLS, observed)),
            "predictions": {"I": dict(zip(M.CELLS, predicted)), "S": dict.fromkeys(M.CELLS, 0.0)}})
    return result


def workbook(path, kind, values=None, duplicate=False, header_error=False):
    book = openpyxl.Workbook()
    sheet = book.active
    identifier = "chr1:100:A:C"
    if kind == "dna":
        sheet.title = "S2.Variant Oligo Seqs"
        headers = {1: "Variant", 2: "Allele 1 Oligo", 3: "Allele 2 Oligo"}
        data = {1: identifier, 2: "A" * 200, 3: "A" * 99 + "C" + "A" * 100}
    else:
        sheet.title = "S4.MPRA results"
        headers = {1: "Variant", 40: "log2Skew in HEPG2", 41: "log2Skew in K562"}
        data = {1: identifier, 40: 1.0, 41: -2.0}
    if header_error:
        headers[1] = "Unexpected ID"
    for col, value in headers.items():
        sheet.cell(2, col, value)
    data.update(values or {})
    for col, value in data.items():
        sheet.cell(3, col, value)
        if duplicate:
            sheet.cell(4, col, value)
    book.save(path)
    book.close()
    return {"sha256": M.sha256(path), "bytes": Path(path).stat().st_size}


class IndependentTransferTests(unittest.TestCase):
    def test_exact_feature_counts_rc_and_delta(self):
        reference = "ACGT" * 50
        self.assertEqual(M.exact_counts(reference).sum(), 298)
        np.testing.assert_array_equal(M.exact_counts(reference), M.exact_counts(M.rc(reference)))
        alternate = reference[:99] + "A" + reference[100:]
        delta = M.exact_counts(alternate) - M.exact_counts(reference)
        self.assertEqual(delta[:-1].sum(), 0)
        self.assertTrue(np.any(delta))
        for bad in ["A" * 199, "a" * 200, "N" * 200]:
            with self.assertRaises(ValueError):
                M.exact_counts(bad)

    def test_zero_feature_change_is_exact(self):
        reference = "A" * 97 + "AAATT" + "C" * 98
        alternate = reference[:99] + "T" + reference[100:]
        np.testing.assert_array_equal(M.exact_counts(reference), M.exact_counts(alternate))

    def test_numpy_five_fold_prediction_and_contexts(self):
        reference = "A" * 200
        alternate = reference[:99] + "C" + reference[100:]
        data = {"x": {"delta_counts": (M.exact_counts(alternate) - M.exact_counts(reference)).tolist()}}
        slopes = {model: {cell: np.zeros((5, 33)) for cell in M.CELLS} for model in ("I", "S")}
        for model in slopes:
            for cell in M.CELLS:
                slopes[model][cell][:, -1] = np.arange(5)
        names, predictions, folds = M.predict_fixed(data, slopes)
        self.assertEqual(names, ["x"])
        for model in slopes:
            for cell in M.CELLS:
                self.assertAlmostEqual(predictions[model][cell][0], 0.01)
                np.testing.assert_array_equal(folds[model][cell][0], np.arange(5) * 0.005)

    def test_macro_component_weighting_has_hand_computed_value(self):
        result = M.evaluate(rows(), replicates=8, seed=0)
        macro = result["observed"]["equal_component"]
        self.assertEqual(macro["losses"]["paired"]["zero"]["mse"], 5.5)
        self.assertEqual(macro["losses"]["paired"]["I"]["mse"], 1.0625)
        self.assertEqual(macro["gains_vs_zero"]["paired"]["I"]["mse"], 4.4375)
        self.assertAlmostEqual(result["observed"]["equal_variant"]["losses"]["paired"]["I"]["mse"],
                               (0.625 + 2.5 + 0.5) / 3)
        self.assertFalse(result["primary"]["minimum_support_passed"])

    def test_macro_is_invariant_to_whole_component_replication(self):
        original = rows()
        replicated = original + [copy.deepcopy(row) | {"id": "z" + row["id"]} for row in original[1:]]
        a = M.evaluate(original, replicates=3)
        b = M.evaluate(replicated, replicates=3)
        self.assertEqual(a["observed"]["equal_component"], b["observed"]["equal_component"])
        self.assertNotEqual(a["observed"]["equal_variant"], b["observed"]["equal_variant"])

    def test_rng_known_splitmix64_values_and_draw_resource_bound(self):
        indices = M.splitmix_indices(13, replicates=1, seed=0)
        expected = [0xE220A8397B1DCDAF, 0x6E789E6AA1B965F4, 0x06C45D188009454F]
        self.assertEqual(indices[0, :3].tolist(), [value % 13 for value in expected])
        with self.assertRaises(ValueError):
            M.splitmix_indices(25001, replicates=2000)

    def test_type7_interpolation(self):
        self.assertEqual(M.type7([0, 10, 20, 30], [0.25, 0.5, 0.75]), [7.5, 15.0, 22.5])

    def test_all_draws_and_stats_are_retained(self):
        result = M.evaluate(rows(), replicates=12, seed=5)
        bootstrap = result["resampling"]
        self.assertEqual(len(bootstrap["resampling_indices"]), 12)
        self.assertEqual(len(bootstrap["sampled_statistics"]), 12)
        for draw, record in zip(bootstrap["resampling_indices"], bootstrap["sampled_statistics"]):
            losses = [result["per_component"][i]["losses"]["paired"]["I"]["mse"] for i in draw]
            self.assertEqual(np.mean(losses), record["equal_component"]["losses"]["paired"]["I"]["mse"])

    def test_association_retains_zeros_and_constant_vector_is_nonestimable(self):
        same = M.association([-1.0, 0.0, 1.0], [-2.0, 0.0, 2.0])
        self.assertEqual(same["sign_agreement_count"], 3)
        self.assertEqual(same["exact_zero_predictions"], 1)
        constant = M.association([0.1] * 7, [0.0] * 7)
        self.assertFalse(constant["pearson"]["estimable"])
        self.assertEqual(constant["sign_agreement_count"], 0)

    def test_gc_subgroups_preserve_and_list_original_components(self):
        result = M.evaluate(rows(), replicates=3)
        a = result["gc_diagnostics_descriptive_only"]["GC_preserving"]
        self.assertEqual(a["represented_component_ids"], ["a"])
        self.assertEqual(a["empty_original_component_ids"], ["b"])
        self.assertEqual(a["variants"], 1)

    def test_source_dna_original_xlsx_and_snv_validation(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S2.xlsx"
            commitment = workbook(path, "dna")
            found, report = M.read_s2_dna(path, commitment, {"chr1:100:A:C"})
            self.assertEqual(report["required_ids"], 1)
            self.assertFalse(found["chr1:100:A:C"]["gc_preserving"])
            self.assertEqual(found["chr1:100:A:C"]["delta_counts"][-1], 1)
            commitment = workbook(path, "dna", values={3: "A" * 98 + "C" + "A" * 101})
            with self.assertRaises(ValueError):
                M.read_s2_dna(path, commitment, {"chr1:100:A:C"})

    def test_source_header_hash_and_duplicate_ids_are_strict(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S2.xlsx"
            commitment = workbook(path, "dna")
            with self.assertRaises(ValueError):
                M.read_s2_dna(path, commitment | {"sha256": "0" * 64}, {"chr1:100:A:C"})
            commitment = workbook(path, "dna", header_error=True)
            with self.assertRaises(ValueError):
                M.read_s2_dna(path, commitment, {"chr1:100:A:C"})
            commitment = workbook(path, "dna", duplicate=True)
            with self.assertRaises(ValueError):
                M.read_s2_dna(path, commitment, {"chr1:100:A:C"})

    def test_source_effects_exact_bits_and_bound_missing_policy(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S4.xlsx"
            commitment = workbook(path, "effects")
            found, report = M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, [])
            self.assertEqual(found["chr1:100:A:C"]["observed"]["K562"], -2.0)
            self.assertEqual(found["chr1:100:A:C"]["observed_float64_little_endian"]["HepG2"],
                             M.float_bits(1.0))
            commitment = workbook(path, "effects", values={40: "NA"})
            with self.assertRaises(ValueError):
                M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, [])
            found, report = M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, ["NA"])
            self.assertEqual(found, {})
            self.assertEqual(report["incomplete"]["chr1:100:A:C"]["missing_contexts"], ["HepG2"])

    def test_source_formula_error_bool_and_numeric_strings_fail(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S4.xlsx"
            for invalid in ["=1+1", "#DIV/0!", True, "1.5", None]:
                commitment = workbook(path, "effects", values={40: invalid})
                with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                    M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, [])

    def test_failed_source_cell_is_retained_in_independent_type_ledger(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S4.xlsx"
            commitment = workbook(path, "effects", values={40: "=1+1"})
            ledger = []
            with self.assertRaises(ValueError):
                M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, [], cell_ledger=ledger)
            self.assertEqual(ledger[-1]["context"], "HepG2")
            self.assertEqual(ledger[-1]["cell_data_type"], "f")
            self.assertEqual(ledger[-1]["status"], "gate_failure")

    def test_source_gate_fails_before_any_workbook_path_is_accessed(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "binding.json"
            M.write_json(path, {"schema_version": 1, "protocol_id": "variant-transfer-v1",
                "protocol_sha256": M.PROTOCOL_SHA, "source_cohort_gate": {"passed": False}})
            args = types.SimpleNamespace(binding=path, expected_binding_sha256=M.sha256(path))
            # No protocol, workbook or model paths even exist on this object.
            with self.assertRaisesRegex(ValueError, "sealed source/cohort gate"):
                M.validate_empirical_binding(args)

    def test_v1_policy_cannot_be_reinterpreted_after_failure(self):
        protocol = {"protocol_id": "variant-transfer-v1"}
        M.check_missing_policy(protocol, {"missing_policy": {"allowed_tokens": []}})
        with self.assertRaisesRegex(ValueError, "v1 remains strict"):
            M.check_missing_policy(protocol, {"missing_policy": {"allowed_tokens": ["NA"]}})

    def test_v2_requires_explicit_exposure_and_unknown_operational_semantics(self):
        protocol = {"protocol_id": "variant-transfer-v2",
            "analysis_classification": "exploratory_after_v1_source_failure",
            "parent_protocol_sha256": M.PROTOCOL_SHA,
            "operational_missingness": copy.deepcopy(M.V2_OPERATIONAL_MISSINGNESS)}
        binding = {"missing_policy": {"allowed_tokens": ["NA"]},
            "operational_missingness": copy.deepcopy(M.V2_OPERATIONAL_MISSINGNESS),
            "exposure": {"v1_failure_observed": True}}
        M.check_missing_policy(protocol, binding)
        for mutation in ("exposure", "semantic_claim", "classification", "parent"):
            p, b = copy.deepcopy(protocol), copy.deepcopy(binding)
            if mutation == "exposure":
                b["exposure"]["v1_failure_observed"] = False
            elif mutation == "semantic_claim":
                b["operational_missingness"]["source_semantics_established"] = True
            elif mutation == "classification":
                p["analysis_classification"] = "outcome_naive"
            else:
                p["parent_protocol_sha256"] = "0" * 64
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                M.check_missing_policy(p, b)

    def test_operational_NA_is_exact_and_accepts_no_other_missing_aliases(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S4.xlsx"
            for token in ["NA ", "na", "N/A", ""]:
                commitment = workbook(path, "effects", values={40: token})
                with self.subTest(token=token), self.assertRaises(ValueError):
                    M.read_s4_effects(path, commitment, {"chr1:100:A:C"}, ["NA"])

    def test_v2_exposure_receipt_preserves_both_parser_counts_and_v1_failure(self):
        receipt = {"protocol_id": "variant-transfer-v2", "protocol_sha256": M.V2_PROTOCOL_SHA,
            "binding_sha256": "a" * 64, "analysis_classification": "exploratory_after_v1_source_failure",
            "v1_primary_selected_numeric_cells_before_failure": 240,
            "v1_independent_selected_numeric_cells_before_failure": 240,
            "v1_primary_selected_effect_cell_ledger_entries": 241,
            "v1_independent_selected_effect_cell_ledger_entries": 241,
            "v1_predictive_performance_evaluated": False, "v1_bootstrap_computed": False,
            "v1_primary_predictions_computed": False, "v1_independent_fixed_predictions_computed": True,
            "v1_comparison_sha256": "a13b8ea6aaf77a473ee50898e0d35031324af5bba6932a9a554d5ab013fc38fc"}
        M.check_second_freeze_exposure(receipt, "a" * 64)
        for field, value in [("v1_independent_selected_numeric_cells_before_failure", 0),
            ("v1_primary_selected_effect_cell_ledger_entries", True),
            ("v1_independent_fixed_predictions_computed", False),
            ("v1_predictive_performance_evaluated", True)]:
            with self.subTest(field=field), self.assertRaises(ValueError):
                M.check_second_freeze_exposure(receipt | {field: value}, "a" * 64)

    def test_second_freeze_requires_committed_exact_bytes(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            frozen = root / "frozen.txt"
            frozen.write_text("metadata only\n")
            subprocess.run(["git", "init", "-q", str(root)], check=True)
            subprocess.run(["git", "-C", str(root), "add", "frozen.txt"], check=True)
            subprocess.run(["git", "-C", str(root), "-c", "user.name=Test", "-c",
                "user.email=test@example.invalid", "commit", "-qm", "freeze"], check=True)
            commit = subprocess.check_output(["git", "-C", str(root), "rev-parse", "HEAD"], text=True).strip()
            M.check_committed_input(root, commit, frozen, M.sha256(frozen))
            with self.assertRaises(ValueError):
                M.check_committed_input(root, commit, frozen, "0" * 64)
            with self.assertRaises(ValueError):
                M.check_committed_input(root, commit[:7], frozen, M.sha256(frozen))

    def test_frozen_metadata_requires_centered_actual_snv_and_gc_boolean(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "cohort.json"
            row = {"id": "chr1:100:A:C", "component": "a", "locus": "locus",
                "sequence_ref": "A" * 200, "sequence_alt": "A" * 99 + "C" + "A" * 100,
                "gc_preserving": False}
            M.write_json(path, {"schema_version": 1, "pairs": [row]})
            binding = {"cohort": {"format": "json", "variant_count": 1}}
            self.assertEqual(set(M.load_frozen_cohort(path, binding)), {row["id"]})
            M.write_json(path, {"schema_version": 1, "pairs": [row | {"gc_preserving": 0}]})
            with self.assertRaises(ValueError):
                M.load_frozen_cohort(path, binding)

    def test_run_manifest_is_exact_and_tamper_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "known.txt").write_text("exact\n")
            M.manifest_directory(root)
            M.verify_run_directory(root)
            (root / "known.txt").write_text("changed\n")
            with self.assertRaises(ValueError):
                M.verify_run_directory(root)

    def test_independent_numpy_reproduces_primary_cli_on_synthetic_handoff(self):
        # The primary is run as a separate executable; no primary implementation
        # is imported or used to derive the independent arrays/statistics.
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            checkout = Path(__file__).resolve().parents[2]
            protocol = json.loads((checkout / "experiments/variant-transfer-v1.json").read_bytes())
            M.write_json(root / "protocol.json", protocol)
            psha = M.sha256(root / "protocol.json")
            M.write_json(root / "binding.json", {"protocol_id": "variant-transfer-v1", "protocol_sha256": psha,
                "source_cohort_gate": {"passed": True, "scope": "synthetic unit fixture only"}})
            bsha = M.sha256(root / "binding.json")
            variants = rows()
            pins = {"schema_version": 1, "protocol_sha256": psha, "binding_sha256": bsha}
            for model in ("I", "S"):
                directory = root / model
                directory.mkdir()
                training = {"artifacts_sha256": protocol["pretrained_model"]["source_artifacts_sha256"],
                            "protocol_sha256": protocol["pretrained_model"]["source_protocol_sha256"]}
                prediction = {"schema_version": 1, "model": model,
                    "endpoint_direction": "alternate minus reference",
                    "commitments": training | {"pairs_sha256": "c" * 64},
                    "pairs": [{"id": r["id"] + "::" + cell, "context": cell,
                        "predicted_alternate_minus_reference": r["predictions"][model][cell]}
                        for r in variants for cell in M.CELLS]}
                M.write_json(directory / "predictions.json", prediction)
                M.write_json(directory / "manifest.json", {"schema_version": 1, "training": training,
                    "pairs_sha256": "c" * 64, "files": {"predictions.json": M.sha256(directory / "predictions.json")}})
                pins["prediction_" + model + "_sha256"] = M.sha256(directory / "predictions.json")
                pins["prediction_" + model + "_manifest_sha256"] = M.sha256(directory / "manifest.json")
            normalized = {"schema_version": 1, "variants": variants} | {key: pins[key] for key in (
                "protocol_sha256", "binding_sha256", "prediction_I_sha256", "prediction_S_sha256")}
            M.write_json(root / "normalized.json", normalized)
            pins["input_sha256"] = M.sha256(root / "normalized.json")
            M.write_json(root / "pins.json", pins)
            subprocess.run([sys.executable, str(checkout / "scripts/evaluate_variant_transfer.py"),
                "--input", str(root / "normalized.json"), "--protocol", str(root / "protocol.json"),
                "--binding", str(root / "binding.json"), "--pins", str(root / "pins.json"),
                "--prediction-I", str(root / "I"), "--prediction-S", str(root / "S"),
                "--out", str(root / "primary")], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
            actual = M.read_json(root / "primary/evaluation.json.gz")
            independent = M.evaluate(variants)
            result = M.compare_trees(actual, independent)
            self.assertTrue(result["matched"], result["failures"][:5])
            self.assertEqual(len(independent["resampling"]["resampling_indices"]), 2000)

    def test_full_original_source_and_integer_numerator_comparison(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            own, observations = root / "own", root / "observations"
            own.mkdir()
            observations.mkdir()
            reference, alternate = "A" * 200, "A" * 99 + "C" + "A" * 100
            identifier = "chr1:100:A:C"
            before = [0] * 33
            after = [0] * 33
            before[M.TRIMERS.index("AAA")] = 198
            after[M.TRIMERS.index("AAA")] = 195
            for word in ("AAC", "ACA", "CAA"):
                after[M.TRIMERS.index(word)] = 1
            after[-1] = 1
            delta = [b - a for a, b in zip(before, after)]
            meta = {"id": identifier, "component": "a", "locus": identifier,
                "gc_preserving": False, "sequence_ref": reference, "sequence_alt": alternate}
            M.write_json(own / "metadata-cohort.json.gz", {identifier: meta})
            M.write_json(own / "source-DNA.json.gz", {identifier: {"reference": reference,
                "alternate": alternate, "reference_counts": before, "alternate_counts": after,
                "delta_counts": delta}})
            values = {"K562": 1.0, "HepG2": -2.0}
            bits = {cell: M.float_bits(value) for cell, value in values.items()}
            M.write_json(own / "source-observations.json.gz", {identifier: {
                "observed": values, "observed_float64_little_endian": bits}})
            M.write_json(own / "source-missingness.json.gz", {"incomplete": {}, "absent_ids": []})
            M.write_json(own / "inputs.json", {"protocol_sha256": M.PROTOCOL_SHA, "binding_sha256": "a" * 64})
            fixed = {"variant_order": [identifier],
                "predictions": {model: {cell: [0.25] for cell in M.CELLS} for model in ("I", "S")},
                "fold_predictions": {model: {cell: [[0.25] * 5] for cell in M.CELLS} for model in ("I", "S")}}
            M.write_json(own / "fixed-predictions.json.gz", fixed)
            M.write_json(observations / "observations.json.gz", {"protocol_sha256": M.PROTOCOL_SHA,
                "binding_sha256": "a" * 64, "variants": [{"id": identifier, "component": "a",
                "locus": identifier, "gc_preserving": False, "reference": reference,
                "alternate": alternate, "complete_case": True, "missing_contexts": [],
                "source_float64_le_hex": bits}]})
            prediction_roots = {}
            for model in ("I", "S"):
                directory = root / model
                directory.mkdir()
                payload = {"model": model, "feature_names": M.FEATURES,
                    "feature_denominators": [198] * 32 + [200],
                    "commitments": {"artifacts_sha256": M.MODEL_SHA},
                    "pairs": [{"id": identifier + "::" + cell, "context": cell, "position_0based": 99,
                        "reference_feature_numerators": before, "alternate_feature_numerators": after,
                        "feature_numerator_delta": delta,
                        "reference_sha256": hashlib.sha256(reference.encode()).hexdigest(),
                        "alternate_sha256": hashlib.sha256(alternate.encode()).hexdigest(),
                        "predicted_alternate_minus_reference": 0.25, "fold_predictions": [0.25] * 5}
                        for cell in M.CELLS]}
                M.write_json(directory / "predictions.json", payload)
                M.manifest_directory(directory)
                prediction_roots[model] = directory
            M.manifest_directory(own)
            M.manifest_directory(observations)
            result = M.compare_raw_sources_and_predictions(own, observations, prediction_roots)
            self.assertTrue(result["matched"], result["failures"])
            self.assertEqual(result["source_float_bits_exact_comparisons"], 2)
            self.assertEqual(result["exact_integer_feature_numerators_compared"], 396)
            self.assertEqual(result["fixed_prediction_numbers_compared"], 24)
            # A coherently rehashed wrong numerator must still fail arithmetic comparison.
            changed = M.read_json(prediction_roots["I"] / "predictions.json")
            changed["pairs"][0]["reference_feature_numerators"][0] -= 1
            M.write_json(prediction_roots["I"] / "predictions.json", changed)
            M.manifest_directory(prediction_roots["I"])
            self.assertFalse(M.compare_raw_sources_and_predictions(own, observations, prediction_roots)["matched"])

    def test_missing_source_id_requires_separately_bound_absence_rule(self):
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "S4.xlsx"
            commitment = workbook(path, "effects")
            with self.assertRaises(ValueError):
                M.read_s4_effects(path, commitment, {"absent"}, [])
            found, report = M.read_s4_effects(path, commitment, {"absent"}, [], absence_allowed=True)
            self.assertEqual(report["absent_ids"], ["absent"])

    def test_comparison_tolerance_and_exact_metadata_failures(self):
        original = M.evaluate(rows(), replicates=3)
        self.assertTrue(M.compare_trees(original, copy.deepcopy(original))["matched"])
        changed = copy.deepcopy(original)
        changed["primary"]["observed_improvement"] += 1e-11
        self.assertTrue(M.compare_trees(original, changed)["matched"])
        changed["primary"]["observed_improvement"] += 1e-8
        self.assertFalse(M.compare_trees(original, changed)["matched"])
        changed = copy.deepcopy(original)
        changed["per_variant"][0]["observed_float64_little_endian"]["K562"] = M.float_bits(1.00000000001)
        self.assertFalse(M.compare_trees(original, changed)["observed_float_bits_and_integer_indices_exact"])
        changed = copy.deepcopy(original)
        changed["resampling"]["resampling_indices"][0][0] = 1 - changed["resampling"]["resampling_indices"][0][0]
        self.assertFalse(M.compare_trees(original, changed)["matched"])
        changed = copy.deepcopy(original)
        changed["variant_count"] = float(changed["variant_count"])
        self.assertFalse(M.compare_trees(original, changed)["matched"])


if __name__ == "__main__":
    unittest.main()

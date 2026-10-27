#!/usr/bin/env python3
"""Independent context-activity-v3 replay; empirical fitting requires a passed source gate."""
from __future__ import annotations

import argparse
from collections import Counter
import datetime
import hashlib
import importlib.util
import json
import math
from pathlib import Path
import platform
import shutil
import sys
import traceback

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BASE_PATH = Path(__file__).with_name("replicate_context_study.py")
SPEC = importlib.util.spec_from_file_location("helix_activity_v3_independent_base", BASE_PATH)
BASE = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = BASE
SPEC.loader.exec_module(BASE)

FROZEN_V3_SHA256 = "64e14037fb1f42ddc1ecaca1d23bb8096fabac1f1e892d105eaa20f525fbbb20"
ENDPOINT = "published_joint_log2_rna_dna_activity"
CONTEXTS = ("HepG2", "K562", "WTC11")
COMPARISONS = (
    ("shared_dna", "additive_context_dna"),
    ("context_mean", "additive_context_dna"),
    ("additive_context_dna", "independent_context_dna"),
)


def checked_protocol(path: Path) -> tuple[dict, str]:
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    if digest != FROZEN_V3_SHA256:
        raise ValueError("Protocol differs from the frozen v3 authority")
    protocol = json.loads(raw)
    if protocol["protocol_id"] != "context-activity-v3" or protocol["status"] != "frozen":
        raise ValueError("Wrong protocol identity/status")
    return protocol, digest


def literal_numeric(cell) -> float:
    """Accept ordinary finite numeric source cells only, never formulas/caches."""
    if cell.data_type != "n" or isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
        raise ValueError("Activity must be an ordinary numeric source cell, not formula/error/text/boolean")
    try:
        value = float(cell.value)
    except (ValueError, OverflowError) as error:
        raise ValueError("Source number is outside finite float64") from error
    if not math.isfinite(value):
        raise ValueError("Activity source value must be finite")
    return value


def read_activity_cells(path: Path, *, sheet_name: str, id_header: str,
                        columns: dict[str, str], required_ids: set[str],
                        header_row: int = 1) -> tuple[dict, dict]:
    """Independently parse an explicitly bound worksheet with openpyxl.

    This low-level routine also supports synthetic tests. Production callers
    must first verify frozen bytes, the central source gate and complete old IDs.
    No source outcome transform or averaging is performed.
    """
    if openpyxl.__version__ != "3.1.5":
        raise ValueError("Expected the pinned openpyxl 3.1.5 parser")
    if not required_ids or any(not isinstance(x, str) or not x for x in required_ids):
        raise ValueError("Required source IDs must be nonempty strings")
    if isinstance(header_row, bool) or not isinstance(header_row, int) or header_row < 1:
        raise ValueError("Header row must be a positive one-based integer")
    if tuple(columns) != CONTEXTS or len(set(columns.values())) != 3:
        raise ValueError("Require the three declared contexts and distinct exact headers")
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError("Exact declared activity worksheet is absent")
        sheet = workbook[sheet_name]
        header = next(sheet.iter_rows(min_row=header_row, max_row=header_row), ())
        values = [cell.value for cell in header]
        required_headers = [id_header, *columns.values()]
        positions = {}
        for name in required_headers:
            matches = [index for index, value in enumerate(values) if value == name]
            if len(matches) != 1:
                raise ValueError(f"Required exact header is absent or duplicated: {name}")
            positions[name] = matches[0]
        header_binding = {name: index + 1 for name, index in positions.items()}
        sheet.reset_dimensions()
        maximum_column = max(positions.values()) + 1
        rows = {}
        seen = Counter()
        malformed_required = []
        nonempty_ids = 0
        blank_ids = 0
        for number, cells in enumerate(
            sheet.iter_rows(min_row=header_row + 1, max_col=maximum_column),
            start=header_row + 1,
        ):
            identifier_cell = cells[positions[id_header]]
            identifier = identifier_cell.value
            if identifier is None or identifier == "":
                blank_ids += 1
                continue
            if not isinstance(identifier, str) or identifier_cell.data_type in {"f", "e", "b"}:
                raise ValueError("Nonempty source ID must be a literal exact string")
            nonempty_ids += 1
            seen[identifier] += 1
            if identifier not in required_ids:
                continue
            if seen[identifier] != 1:
                raise ValueError(f"Duplicate required source ID: {identifier}")
            try:
                observed = {context: literal_numeric(cells[positions[name]])
                            for context, name in columns.items()}
            except ValueError as error:
                malformed_required.append({"id": identifier, "source_row_1based": number, "reason": str(error)})
                continue
            rows[identifier] = {"source_row_1based": number, "observations": observed}
        missing = sorted(required_ids - set(rows))
        report = {
            "parser": "independent openpyxl read_only/data_only=False; no root importer",
            "openpyxl_version": openpyxl.__version__,
            "worksheet": sheet_name,
            "header_row_1based": header_row,
            "header_binding_1based": header_binding,
            "required_ID_count": len(required_ids),
            "finite_complete_required_ID_count": len(rows),
            "source_nonempty_ID_rows": nonempty_ids,
            "source_unique_IDs": len(seen),
            "source_blank_ID_rows": blank_ids,
            "row_count_semantics": "Nonempty IDs counted explicitly; openpyxl fills worksheet row gaps, so blank-ID count includes gaps and is not a physical XML row count.",
            "extra_source_IDs": sorted(set(seen) - required_ids),
            "duplicate_extra_source_IDs": {key: value for key, value in sorted(seen.items()) if value > 1 and key not in required_ids},
            "invalid_required_rows": malformed_required,
            "missing_required_IDs": missing,
            "activity_transform": "none; source float values preserved unchanged",
        }
        if malformed_required or missing:
            raise ValueError("Required activity coverage is incomplete or nonnumeric: " + json.dumps(report))
        return rows, report
    finally:
        workbook.close()


def activity_evaluation(records: list[dict], contexts, folds, alpha=10.0) -> dict:
    """Reuse inspected independent fits, replacing only endpoint metadata."""
    result = BASE.evaluate_complete_panel(records, contexts, folds, alpha)
    result["endpoint"] = ENDPOINT
    result["fixed_zero_prediction_audit"]["scope"] = "Published uncentered activity; fixed descriptive zero predictor, no outcome transform"
    result["common_activity"] = common_activity_diagnostic(result)
    return result


def common_activity_diagnostic(evaluation: dict) -> dict:
    """Independently check the C-vs-A / G-vs-S mean-activity loss identity."""
    contexts = tuple(evaluation["contexts"])
    rows = {}
    for prediction in evaluation["predictions"]:
        identifier = prediction["sequence_id"]
        cell = prediction["context"]
        if cell not in contexts or cell in rows.setdefault(identifier, {}):
            raise ValueError("Incomplete/duplicate context prediction")
        rows[identifier][cell] = prediction
    details = []
    maximum_residual = 0.0
    for identifier in sorted(rows):
        by_context = rows[identifier]
        if set(by_context) != set(contexts):
            raise ValueError("Common activity requires a complete context triple")
        observed = np.array([by_context[c]["observed"] for c in contexts], dtype=float)
        global_prediction = np.array([by_context[c]["predictions"]["global_mean"] for c in contexts])
        shared_prediction = np.array([by_context[c]["predictions"]["shared_dna"] for c in contexts])
        if not np.array_equal(global_prediction, np.repeat(global_prediction[0], len(contexts))):
            raise ValueError("Global predictor unexpectedly varies across contexts")
        if not np.array_equal(shared_prediction, np.repeat(shared_prediction[0], len(contexts))):
            raise ValueError("Shared predictor unexpectedly varies across contexts")
        context_prediction = np.array([by_context[c]["predictions"]["context_mean"] for c in contexts])
        additive_prediction = np.array([by_context[c]["predictions"]["additive_context_dna"] for c in contexts])
        actual_mean = float(observed.mean())
        mean_global_mse = (actual_mean - global_prediction[0]) ** 2
        mean_shared_mse = (actual_mean - shared_prediction[0]) ** 2
        primary_gain = float(np.mean((observed - context_prediction) ** 2) -
                             np.mean((observed - additive_prediction) ** 2))
        mean_gain = float(mean_global_mse - mean_shared_mse)
        residual = primary_gain - mean_gain
        maximum_residual = max(maximum_residual, abs(residual))
        exemplar = by_context[contexts[0]]
        details.append({"sequence_id": identifier, "group": exemplar["group"],
                        "fold_index_0based": exemplar["fold_index_0based"],
                        "observed_mean": actual_mean,
                        "global_mean_prediction": float(global_prediction[0]),
                        "shared_dna_prediction": float(shared_prediction[0]),
                        "global_mean_squared_error": float(mean_global_mse),
                        "shared_dna_squared_error": float(mean_shared_mse),
                        "C_minus_A_cell_MSE": primary_gain,
                        "G_minus_S_mean_activity_MSE": mean_gain,
                        "identity_residual": residual})
    if not details:
        raise ValueError("Common activity diagnostic is empty")
    return {
        "scope": "Correlated algebraic diagnostic, not independent biological validation",
        "records": details,
        "global_mean_mse": float(np.mean([x["global_mean_squared_error"] for x in details])),
        "shared_dna_mse": float(np.mean([x["shared_dna_squared_error"] for x in details])),
        "mean_activity_mse_improvement": float(np.mean([x["G_minus_S_mean_activity_MSE"] for x in details])),
        "primary_cell_mse_improvement": float(np.mean([x["C_minus_A_cell_MSE"] for x in details])),
        "maximum_absolute_per_element_identity_residual": maximum_residual,
    }


def verified_dataset_gate(directory: Path, protocol: dict, protocol_sha: str) -> tuple[dict, dict, Path]:
    manifest_path = directory / "manifest.json"
    manifest = json.loads(manifest_path.read_bytes())
    if manifest.get("schema_version") != 1 or manifest.get("protocol_sha256") != protocol_sha or manifest.get("endpoint") != ENDPOINT:
        raise ValueError("Activity dataset manifest protocol/endpoint/schema mismatch")
    required_gates = (
        "archive_integrity_verified", "empirical_activity_schema_bound", "coverage_complete",
        "design_identity_verified", "replicate_aggregation_status_recorded",
    )
    if any(manifest.get("gates", {}).get(name) is not True for name in required_gates):
        raise ValueError("The central activity source gate has not passed")
    if not manifest.get("aggregation_status") or not manifest.get("schema_binding"):
        raise ValueError("Source aggregation uncertainty and schema binding must be recorded")
    required_files = {
        "activity.json.gz", "source-binding.json", "source-export.json.gz", "source-ledger.jsonl.gz",
        "recovery-summary.json", "archive-validation.json", "extraction-receipt.json",
    }
    if set(manifest.get("files", {})) != required_files:
        raise ValueError("Activity source receipt inventory differs")
    for name, expected in manifest["files"].items():
        if BASE.file_hashes(directory / name)["sha256"] != expected:
            raise ValueError("Activity dataset commitment mismatch: " + name)
    source = manifest["source"]
    path = (directory / source["raw_path"]).resolve()
    observed = BASE.file_hashes(path)
    expected_sha = protocol["source"]["jointdata_expected_sha256"]
    expected_size = protocol["source"]["expected_bytes"]
    if source["sha256"] != expected_sha or source["bytes"] != expected_size:
        raise ValueError("Dataset gate refers to another activity workbook")
    if observed["sha256"] != expected_sha or observed["bytes"] != expected_size:
        raise ValueError("Frozen raw activity workbook hash/size mismatch")
    return manifest, {"manifest_sha256": BASE.file_hashes(manifest_path)["sha256"],
                      "source_workbook_sha256": observed["sha256"], "source_workbook_bytes": observed["bytes"],
                      "source_gate_files": manifest["files"]}, path


def checked_metadata(protocol: dict) -> tuple[list[dict], dict, dict]:
    panel_path = ROOT / protocol["panel"]["path"]
    groups_path = ROOT / protocol["grouping"]["path"]
    if BASE.file_hashes(panel_path)["sha256"] != protocol["panel"]["sha256"]:
        raise ValueError("Prior source DNA panel bytes differ")
    if BASE.file_hashes(groups_path)["sha256"] != protocol["grouping"]["sha256"]:
        raise ValueError("Frozen prior group/fold bytes differ")
    source = json.loads(panel_path.read_bytes())
    groups = json.loads(groups_path.read_bytes())
    # Construct metadata-only records: old ESS observations are deliberately discarded.
    records = [{"id": row["id"], "sequence": row["sequence"]} for row in source["records"]]
    records.sort(key=lambda row: row["id"])
    identifiers = [row["id"] for row in records]
    if len(records) != protocol["panel"]["elements"] or len(set(identifiers)) != len(records):
        raise ValueError("Prior fixed element inventory differs")
    if any(len(row["sequence"]) != 200 or set(row["sequence"]) - set("ACGT") for row in records):
        raise ValueError("Prior metadata contains an invalid 200-base core")
    mapping = {row["id"]: row["group"] for row in groups["records"]}
    if set(mapping) != set(identifiers) or len(groups["records"]) != len(records):
        raise ValueError("Frozen group IDs differ from prior fixed elements")
    if len(set(mapping.values())) != 2761 or len(groups["folds"]) != 5:
        raise ValueError("Frozen group/fold count differs")
    for row in records:
        row["group"] = mapping[row["id"]]
    return records, groups, {
        "prior_metadata_panel_sha256": protocol["panel"]["sha256"],
        "grouping_sha256": protocol["grouping"]["sha256"],
        "ESS_observations_used": False,
        "DNA_group_reference_scope": "Reuse previously independently validated metadata/groups exactly; re-extract every feature and refit every outcome model.",
    }


def run_replay(protocol_path: Path, dataset: Path, output: Path, source_only: bool = False) -> None:
    if output.exists():
        raise FileExistsError("Independent activity output already exists; retain each attempt")
    output.mkdir(parents=True)
    shutil.copyfile(__file__, output / "implementation.py")
    shutil.copyfile(BASE_PATH, output / "independent-context-base.py")
    started = datetime.datetime.now(datetime.timezone.utc).isoformat()
    try:
        protocol, digest = checked_protocol(protocol_path)
        gate, gate_commitments, workbook_path = verified_dataset_gate(dataset, protocol, digest)
        records, grouping, metadata_commitments = checked_metadata(protocol)
        BASE.write_json(output / "inputs.json", {
            "protocol_sha256": digest, "endpoint": ENDPOINT, **gate_commitments, **metadata_commitments,
            "source": gate["source"], "aggregation_status": gate["aggregation_status"],
            "schema_binding": gate["schema_binding"], "started_at_utc": started,
            "implementation_sha256": BASE.file_hashes(Path(__file__))["sha256"],
            "independent_base_sha256": BASE.file_hashes(BASE_PATH)["sha256"],
            "python": platform.python_version(), "numpy": np.__version__, "openpyxl": openpyxl.__version__,
            "dependency_lock_sha256": BASE.file_hashes(ROOT / "requirements/research.lock")["sha256"],
        })
        print("Independently parsing the authorized raw activity workbook with openpyxl", flush=True)
        measured, source_report = read_activity_cells(
            workbook_path, sheet_name=protocol["source"]["worksheet"],
            id_header=protocol["source"]["id_header"], columns=protocol["source"]["columns"],
            required_ids={row["id"] for row in records}, header_row=1)
        if len(measured) != 55276:
            raise ValueError("All 55,276 fixed IDs are required; subset selection is forbidden")
        for row in records:
            row["observations"] = measured[row["id"]]["observations"]
            row["source_row_1based"] = measured[row["id"]]["source_row_1based"]
        BASE.write_json(output / "source-verification.json", source_report)
        BASE.write_json(output / "source-records.json.gz", records)
        BASE.write_json(output / "group-input.json", {
            "path": protocol["grouping"]["path"], "sha256": protocol["grouping"]["sha256"]})
        if not source_only:
            print("Fitting all v3 activity models independently; no old outcome fits/caches", flush=True)
            evaluation = activity_evaluation(records, CONTEXTS, grouping["folds"], protocol["models"]["alpha"])
            evaluation["protocol_sha256"] = digest
            BASE.write_json(output / "evaluation.json.gz", evaluation)
            BASE.write_json(output / "common-activity.json.gz", evaluation["common_activity"])
            bootstrap = BASE.paired_component_bootstrap(
                [row["group"] for row in records],
                [row["model_losses"]["mse"] for row in evaluation["per_sequence_errors"]],
                [row["model_losses"]["mae"] for row in evaluation["per_sequence_errors"]],
                COMPARISONS, replicates=2000, seed=0x2026102702)
            BASE.write_json(output / "resampling.json.gz", bootstrap)
            summary = {key: evaluation[key] for key in (
                "schema_version", "endpoint", "protocol_sha256", "contexts", "model_order",
                "sequence_count", "group_count", "measurement_count", "models", "balanced_design_checks",
            )}
            summary["primary_index"] = 1
            summary["planned_comparisons"] = bootstrap["equal_sequences"]["comparisons"]
            summary["common_activity"] = {key: value for key, value in evaluation["common_activity"].items() if key != "records"}
            BASE.write_json(output / "summary.json", summary)
        BASE.write_json(output / "status.json", {
            "status": "source_verified" if source_only else "complete",
            "protocol_sha256": digest, "endpoint": ENDPOINT,
            "completed_at_utc": datetime.datetime.now(datetime.timezone.utc).isoformat(),
            "elements": len(records), "observations": len(records) * len(CONTEXTS),
        })
    except BaseException as error:
        BASE.write_json(output / "failure.json", {
            "status": "failed_retained_attempt", "error_type": type(error).__name__, "error": str(error),
            "traceback": traceback.format_exc(), "started_at_utc": started,
            "model_fits_exported": (output / "evaluation.json.gz").exists(),
        })
        raise
    finally:
        BASE.write_json(output / "manifest.json", {
            "schema_version": 1, "protocol_sha256": FROZEN_V3_SHA256, "endpoint": ENDPOINT,
            "files": {path.name: BASE.file_hashes(path)["sha256"] for path in sorted(output.iterdir())
                      if path.is_file() and path.name != "manifest.json"},
        })


def compare_rust_run(independent_root: Path, rust_study_path: Path, output: Path) -> bool:
    if output.exists():
        raise FileExistsError("Activity comparison output already exists; retain every attempt")
    output.mkdir(parents=True)
    shutil.copyfile(__file__, output / "implementation.py")
    shutil.copyfile(BASE_PATH, output / "independent-context-base.py")
    try:
        _, own_commitment = BASE.verified_output_manifest(independent_root, "manifest.json", 1)
        root = rust_study_path.resolve().parent
        manifest, rust_commitment = BASE.verified_output_manifest(root, "artifacts.json", 5)
        required = {"features.json.gz", "study.json.gz", "resampling.json.gz", "summary.json",
                    "protocol.json", "inputs.json", "grouping.json", "common-activity.json.gz"}
        if not required.issubset(manifest["files"]) or rust_study_path.resolve() != root / "study.json.gz":
            raise ValueError("Rust schema5 core inventory/study path differs")
        protocol, _ = checked_protocol(root / "protocol.json")
        if BASE.file_hashes(root / "grouping.json")["sha256"] != protocol["grouping"]["sha256"]:
            raise ValueError("Rust core does not retain exact frozen group bytes")
        own_inputs = json.loads((independent_root / "inputs.json").read_bytes())
        if own_inputs["protocol_sha256"] != FROZEN_V3_SHA256 or own_inputs["source_workbook_sha256"] != protocol["source"]["jointdata_expected_sha256"]:
            raise ValueError("Independent activity source/protocol commitment mismatch")
        own, own_study_receipt = BASE.read_json_payload(independent_root / "evaluation.json.gz")
        rust, rust_study_receipt = BASE.read_json_payload(rust_study_path)
        if own["endpoint"] != ENDPOINT:
            raise ValueError("Independent replay endpoint differs")
        own_boot, own_boot_receipt = BASE.read_json_payload(independent_root / "resampling.json.gz")
        rust_boot, rust_boot_receipt = BASE.read_json_payload(root / "resampling.json.gz")
        rust_features, feature_receipt = BASE.read_json_payload(root / "features.json.gz")
        records, source_receipt = BASE.read_json_payload(independent_root / "source-records.json.gz")
        common, common_receipt = BASE.read_json_payload(root / "common-activity.json.gz")
        comparison = BASE.NumericalComparison()
        BASE.compare_sequence_features(own, records, rust_features, rust, comparison)
        BASE.compare_evaluations(own, rust, comparison)
        BASE.compare_bootstraps(own_boot, rust_boot, comparison, own["models"])
        own_common = own["common_activity"]
        comparison.exact("activity_common_ID_order", [row["sequence_id"] for row in own_common["records"]],
                         [row["sequence_id"] for row in common["records"]])
        comparison.exact("activity_common_group_order", [row["group"] for row in own_common["records"]],
                         [row["group"] for row in common["records"]])
        comparison.exact("activity_common_fold_order", [row["fold_index_0based"] for row in own_common["records"]],
                         [row["fold_index_0based"] for row in common["records"]])
        for key in ("observed_mean", "global_mean_prediction", "shared_dna_prediction",
                    "global_mean_squared_error", "shared_dna_squared_error",
                    "C_minus_A_cell_MSE", "G_minus_S_mean_activity_MSE", "identity_residual"):
            tolerance = 1e-7 if key in {"observed_mean", "global_mean_prediction", "shared_dna_prediction"} else 1e-8
            comparison.numeric("activity_common_" + key, [row[key] for row in own_common["records"]],
                               [row[key] for row in common["records"]], tolerance)
        for key in ("global_mean_mse", "shared_dna_mse", "mean_activity_mse_improvement",
                    "primary_cell_mse_improvement", "maximum_absolute_per_element_identity_residual"):
            comparison.numeric("activity_common_summary_" + key, own_common[key], common[key], 1e-8)
        comparison.exact("activity_bootstrap_primary_identity",
                         own_boot["comparison_order"][1],
                         {"baseline": "context_mean", "candidate": "additive_context_dna"})
        result = comparison.report()
        result.update({
            "schema_version": 1, "protocol_sha256": FROZEN_V3_SHA256, "endpoint": ENDPOINT,
            "independent_manifest": own_commitment, "rust_artifacts_manifest": rust_commitment,
            "independent_study": own_study_receipt, "rust_study": rust_study_receipt,
            "independent_resampling": own_boot_receipt, "rust_resampling": rust_boot_receipt,
            "rust_features": feature_receipt, "independent_raw_source_records": source_receipt,
            "rust_common_activity": common_receipt,
            "implementation_sha256": BASE.file_hashes(Path(__file__))["sha256"],
            "independent_base_sha256": BASE.file_hashes(BASE_PATH)["sha256"],
            "prediction_absolute_tolerance": 1e-7, "metric_absolute_tolerance": 1e-8,
            "limits": ["Same source experiment and source-selected panel; numerical agreement is not independent biological replication.",
                       "Source-normalized published log2 RNA/DNA is not absolute transcripts per cell.",
                       "The common-component identity is an algebraic diagnostic, not independent biological corroboration."],
        })
        BASE.write_json(output / "comparison.json", result)
        print(json.dumps({"matched": result["matched"], "failure_count": result["failure_count"],
                          "comparison": str(output / "comparison.json")}))
        return result["matched"]
    except BaseException as error:
        BASE.write_json(output / "failure.json", {"status": "failed_retained_comparison",
                        "error_type": type(error).__name__, "error": str(error), "traceback": traceback.format_exc()})
        raise
    finally:
        BASE.write_json(output / "manifest.json", {
            "schema_version": 1, "protocol_sha256": FROZEN_V3_SHA256, "endpoint": ENDPOINT,
            "files": {path.name: BASE.file_hashes(path)["sha256"] for path in sorted(output.iterdir())
                      if path.is_file() and path.name != "manifest.json"},
        })


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--protocol", type=Path, default=ROOT / "experiments/context-activity-v3.json")
    parser.add_argument("--dataset", type=Path, default=ROOT / "datasets/context-activity")
    parser.add_argument("--out", type=Path)
    parser.add_argument("--source-check-only", action="store_true")
    parser.add_argument("--compare-rust", type=Path)
    parser.add_argument("--independent-run", type=Path)
    parser.add_argument("--describe", action="store_true")
    args = parser.parse_args()
    if args.describe:
        print(json.dumps({"protocol_sha256": FROZEN_V3_SHA256, "endpoint": ENDPOINT,
                          "source_gate_required": True, "empirical_data_loaded": False,
                          "implementation_sha256": BASE.file_hashes(Path(__file__))["sha256"]}))
        return
    if args.out is None:
        parser.error("--out NEW_DIRECTORY is required")
    if args.compare_rust:
        if args.independent_run is None:
            parser.error("--compare-rust requires --independent-run")
        raise SystemExit(0 if compare_rust_run(args.independent_run, args.compare_rust, args.out) else 2)
    run_replay(args.protocol, args.dataset, args.out, source_only=args.source_check_only)


if __name__ == "__main__":
    main()

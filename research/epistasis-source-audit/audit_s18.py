#!/usr/bin/env python3
"""Design-only audit; reads S18 pair oligos, never S20 response values."""
import argparse
import collections
import hashlib
import json
import pathlib
import re

import openpyxl

SOURCE_SHA256 = "92929b4cef82797237c4c2be770feb744daa97255ac6dbfe3d53d5a34c996fbb"
HEADERS = ("Variant 1", "Variant 2", "Distance", "Center Variant", "Window", "Diplotype", "Oligo")
LABELS = ("RR", "AR", "RA", "AA")
VARIANT = re.compile(r"^(chr(?:[1-9]|1[0-9]|2[0-2]|X|Y)):(\d+):([ACGT]):([ACGT])$")


def changes(a, b):
    return [(i, x, y) for i, (x, y) in enumerate(zip(a, b)) if x != y]


def audit(path):
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    if digest != SOURCE_SHA256:
        raise ValueError("unexpected S18 source hash")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if wb.sheetnames != ["S18.Variant Pair Oligos"]:
        raise ValueError("unexpected sheets")
    rows = wb.active.iter_rows(values_only=True)
    title, headers = next(rows), next(rows)
    if tuple(headers) != HEADERS:
        raise ValueError("unexpected header")
    groups = collections.defaultdict(list)
    sequence_lengths = collections.Counter()
    non_acgt_rows = 0
    for n, row in enumerate(rows, start=3):
        if len(row) != 7 or any(isinstance(x, str) and x.startswith("=") for x in row):
            raise ValueError(f"formula or malformed row {n}")
        v1, v2, distance, center, window, label, oligo = row
        groups[(v1, v2, center, window)].append((n, distance, label, oligo))
        if isinstance(oligo, str):
            sequence_lengths[str(len(oligo))] += 1
            non_acgt_rows += bool(set(oligo) - set("ACGT"))

    status = collections.Counter()
    label_counts = collections.Counter()
    distance_counts = collections.Counter()
    issues = []
    eligible = []
    unique_pairs = set()
    reference_hashes = collections.defaultdict(list)
    for key, items in sorted(groups.items()):
        v1, v2, center, window = key
        unique_pairs.add((v1, v2))
        labels = collections.Counter(item[2] for item in items)
        label_counts.update(item[2] for item in items)
        reasons = []
        m1, m2 = VARIANT.fullmatch(str(v1)), VARIANT.fullmatch(str(v2))
        if not m1 or not m2 or m1.group(1) != m2.group(1):
            reasons.append("non_snv_or_cross_chromosome_identifier")
        if center not in ("var1", "var2") or window not in ("left", "middle", "right"):
            reasons.append("unexpected_window_label")
        if any(n > 1 for n in labels.values()):
            reasons.append("duplicate_diplotype_label")
        if set(labels) != set(LABELS):
            reasons.append("missing_or_extra_diplotype_label")
        if len(items) != 4:
            reasons.append("not_exactly_four_rows")
        if any(not isinstance(seq, str) or len(seq) != 200 or set(seq) - set("ACGT") for _, _, _, seq in items):
            reasons.append("not_literal_200nt_acgt")
        if m1 and m2 and m1.group(1) == m2.group(1):
            chrom, p1, r1, a1 = m1.groups()
            _, p2, r2, a2 = m2.groups()
            distance = abs(int(p1) - int(p2))
            if any(type(d) is not int or d != distance for _, d, _, _ in items):
                reasons.append("distance_mismatch")
            if r1 == a1 or r2 == a2 or distance == 0:
                reasons.append("invalid_snv_identity")
            if distance < 3:
                reasons.append("distance_below_three")
        if not reasons:
            seqs = {label: seq for _, _, label, seq in items}
            first = changes(seqs["RR"], seqs["AR"])
            second = changes(seqs["RR"], seqs["RA"])
            if len(first) != 1 or len(second) != 1:
                reasons.append("single_labels_not_single_substitutions")
            elif first[0][0] == second[0][0]:
                reasons.append("substitutions_at_same_position")
            else:
                p, ref1, alt1 = first[0]
                q, ref2, alt2 = second[0]
                expected = list(seqs["RR"])
                expected[p], expected[q] = alt1, alt2
                if "".join(expected) != seqs["AA"]:
                    reasons.append("double_not_combination")
                if (ref1, alt1, ref2, alt2) != (r1, a1, r2, a2):
                    reasons.append("labeled_alleles_not_forward_genomic_identifiers")
                if abs(p - q) != distance:
                    reasons.append("sequence_distance_mismatch")
        if reasons:
            for reason in reasons:
                status[reason] += 1
            issues.append({"key": list(key), "rows": [i[0] for i in items], "reasons": reasons})
        else:
            status["eligible"] += 1
            distance_counts[str(distance)] += 1
            eligible.append({"variant1": v1, "variant2": v2, "center": center, "window": window,
                             "distance": distance, "rows": {label: n for n, _, label, _ in items},
                             "position1_zero_based": p, "position2_zero_based": q})
            ref_hash = hashlib.sha256(seqs["RR"].encode("ascii")).hexdigest()
            reference_hashes[ref_hash].append((v1, v2, center, window))
    result = {"source_file": path.name, "source_sha256": digest, "sheet": wb.active.title,
              "title": title[0], "headers": list(headers), "data_rows": sum(len(v) for v in groups.values()),
              "design_groups": len(groups), "distinct_pair_identifiers": len(unique_pairs),
              "diplotype_rows": dict(sorted(label_counts.items())),
              "sequence_length_rows": dict(sorted(sequence_lengths.items(), key=lambda x: int(x[0]))),
              "rows_with_non_acgt_characters": non_acgt_rows,
              "status_counts_not_mutually_exclusive": dict(sorted(status.items())),
              "eligible_distinct_pairs": len({(e['variant1'], e['variant2']) for e in eligible}),
              "eligible_reference_sequence_reuse": {
                  "unique_sha256_fingerprints": len(reference_hashes),
                  "groups_in_nonunique_fingerprints": sum(len(v) for v in reference_hashes.values() if len(v) > 1),
                  "extra_groups_after_one_per_fingerprint": sum(len(v) - 1 for v in reference_hashes.values()),
                  "largest_fingerprint_group": max((len(v) for v in reference_hashes.values()), default=0),
                  "nonunique_fingerprints_spanning_multiple_pair_ids": sum(
                      len(v) > 1 and len({(x[0], x[1]) for x in v}) > 1 for v in reference_hashes.values()),
              },
              "eligible_distance_distribution": dict(sorted(distance_counts.items(), key=lambda x: int(x[0]))),
              "eligible_groups": eligible, "ineligible_groups": issues,
              "scope": "S18 design only; no S20 observations or outcome availability inspected"}
    wb.close()
    return result


def main():
    p = argparse.ArgumentParser()
    p.add_argument("source", type=pathlib.Path)
    p.add_argument("output", type=pathlib.Path)
    args = p.parse_args()
    result = audit(args.source)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")
    print(json.dumps({k: v for k, v in result.items() if k not in ("eligible_groups", "ineligible_groups", "eligible_distance_distribution")}, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()

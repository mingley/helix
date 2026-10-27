"""Synthetic independent-parser and equation tests; no new source payloads."""
import contextlib
import hashlib
import importlib.util
import io
import json
from pathlib import Path
import re
import sys
import tempfile
import unittest
from fractions import Fraction
import warnings
import xml.etree.ElementTree as ET
import zipfile

import openpyxl
import numpy as np


PATH = Path(__file__).resolve().parents[2] / "scripts" / "verify_context_endpoint_recovery.py"
SPEC = importlib.util.spec_from_file_location("independent_endpoint_recovery", PATH)
RECOVERY = importlib.util.module_from_spec(SPEC)
sys.modules[SPEC.name] = RECOVERY
SPEC.loader.exec_module(RECOVERY)


class RecoveryGoldenTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.path = self.root / "synthetic.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "target"
        sheet.append(["ID", "Cell A", "Cell B", "Cell C"])
        sheet.append(["one", 0.1, True, "#N/A"])
        sheet["D2"].data_type = "e"
        sheet["A4"] = " second "
        sheet["B4"] = "=B2+1"
        sheet["A7"].font = openpyxl.styles.Font(bold=True)
        sheet.row_dimensions[9].height = 12
        workbook.save(self.path)
        workbook.close()
        self.patch_sheet()

    def patch_sheet(self):
        with zipfile.ZipFile(self.path) as source:
            contents = {item.filename: source.read(item.filename) for item in source.infolist()}
        name = "xl/worksheets/sheet1.xml"
        root = ET.fromstring(contents[name])
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        root.find(ns + "dimension").set("ref", "A1:XFD1048576")
        for cell in root.findall(".//" + ns + "c"):
            if cell.attrib["r"] == "B2":
                cell.find(ns + "v").text = "0.10000000000000001"
            if cell.attrib["r"] == "B4":
                cell.find(ns + "v").text = "1.1"
            if cell.attrib["r"] == "A1":
                inline = cell.find(ns + "is")
                for child in list(inline):
                    inline.remove(child)
                for letter in ["I", "D"]:
                    run = ET.SubElement(inline, ns + "r")
                    ET.SubElement(run, ns + "t").text = letter
        contents[name] = ET.tostring(root)
        with zipfile.ZipFile(self.path, "w", compression=zipfile.ZIP_DEFLATED) as destination:
            for filename, content in contents.items():
                destination.writestr(filename, content)

    def test_physical_rows_do_not_expand_advisory_million_row_dimension(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            inventory = workbook.inventory({"target": [1]})
        sheet = inventory["sheets"][0]
        self.assertEqual(sheet["declared_dimension_advisory"], {"max_row": 1048576, "max_column": 16384})
        self.assertEqual(sheet["counts"]["physical_stored_rows"], 5)
        self.assertEqual(sheet["counts"]["meaningful_stored_rows"], 3)
        self.assertEqual(sheet["counts"]["explicit_stored_blank_rows"], 2)
        self.assertEqual(sheet["counts"]["unrepresented_gaps_within_stored_extent"], 4)
        self.assertEqual(sheet["largest_stored_row"], 9)

    def test_header_rich_text_and_exact_whitespace_strings_are_preserved(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            rows = dict(workbook.physical_rows("target"))
        self.assertEqual(rows[1][1].value, "ID")
        self.assertEqual(rows[4][1].value, " second ")

    def test_numeric_literal_survives_library_float_conversion_exactly(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            row = dict(workbook.physical_rows("target"))[2]
        self.assertEqual(row[2].value, 0.1)
        self.assertEqual(row[2].raw_v, "0.10000000000000001")
        self.assertEqual(row[2].numeric_fraction(), Fraction(10000000000000001, 100000000000000000))

    def test_formulas_are_counted_and_numeric_cache_needs_explicit_permission(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            rows = dict(workbook.physical_rows("target"))
            inventory = workbook.inventory({"target": [1]})
        formula = rows[4][2]
        self.assertTrue(formula.formula_present)
        self.assertEqual(formula.formula, "B2+1")
        self.assertEqual(formula.raw_v, "1.1")
        self.assertEqual(inventory["sheets"][0]["counts"]["formula_cells_with_cache_literal"], 1)
        with self.assertRaisesRegex(ValueError, "explicit frozen"):
            formula.numeric_fraction()
        self.assertEqual(formula.numeric_fraction(allow_cached_formula=True), Fraction(11, 10))

    def test_booleans_errors_and_style_only_blanks_never_become_zero_measurements(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            rows = dict(workbook.physical_rows("target"))
        for cell in [rows[2][3], rows[2][4], rows[7][1]]:
            with self.subTest(kind=cell.kind), self.assertRaises(ValueError):
                cell.numeric_fraction()
        self.assertEqual(rows[2][3].kind, "b")
        self.assertEqual(rows[2][4].kind, "e")
        self.assertFalse(rows[7][1].meaningful)

    def test_inventory_exposes_only_declared_header_values(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            inventory = workbook.inventory({"target": [1]})
        serialized = json.dumps(inventory)
        self.assertNotIn("0.10000000000000001", serialized)
        self.assertNotIn(" second ", serialized)
        self.assertNotIn("B2+1", serialized)
        self.assertEqual(set(inventory["sheets"][0]["declared_headers_only"]), {"1"})

    def test_absent_sheet_and_unstored_header_row_fail(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            with self.assertRaisesRegex(ValueError, "absent"):
                workbook.inventory({"missing": [1]})
            with self.assertRaisesRegex(ValueError, "not physically stored"):
                workbook.inventory({"target": [3]})

    def test_expected_sha256_is_checked_before_library_parse_or_output_creation(self):
        out = self.root / "output"
        with self.assertRaisesRegex(ValueError, "integrity"):
            RECOVERY.inventory_file(self.path, "0" * 64, {"target": [1]}, out)
        self.assertFalse(out.exists())

    def test_inventory_outputs_commitments_and_never_overwrites(self):
        out = self.root / "output"
        expected = hashlib.sha256(self.path.read_bytes()).hexdigest()
        with contextlib.redirect_stdout(io.StringIO()):
            RECOVERY.inventory_file(self.path, expected, {"target": [1]}, out)
        manifest = json.loads((out / "manifest.json").read_bytes())
        for name, digest in manifest["files"].items():
            self.assertEqual(hashlib.sha256((out / name).read_bytes()).hexdigest(), digest)
        with self.assertRaises(FileExistsError):
            RECOVERY.inventory_file(self.path, expected, {"target": [1]}, out)

    def test_duplicate_zip_entries_fail_before_parsing(self):
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            with zipfile.ZipFile(self.path, "a") as archive:
                archive.writestr("xl/workbook.xml", archive.read("xl/workbook.xml"))
        with self.assertRaisesRegex(ValueError, "duplicate XLSX"):
            RECOVERY.IndependentWorkbook(self.path)

    def bound_table(self, *, allow_formula=False):
        with RECOVERY.IndependentWorkbook(self.path) as workbook:
            return RECOVERY.extract_bound_table(workbook, sheet="target", header_row=1, first_data_row=2,
                id_column=1, value_columns={"a": 2, "b": 3, "c": 4},
                expected_headers={1: "ID", 2: "Cell A", 3: "Cell B", 4: "Cell C"},
                allow_cached_formula=allow_formula)

    def test_bound_extraction_retains_all_missingness_without_boolean_or_error_imputation(self):
        table = self.bound_table()
        self.assertEqual(table["accounting"]["physical_data_rows"], 4)
        self.assertEqual(table["accounting"]["blank_id_rows"], 2)
        self.assertEqual(table["accounting"]["unique_ids"], 2)
        self.assertEqual(table["accounting"]["complete_ids"], 0)
        self.assertEqual(table["records"][0]["values"], {"a": 0.1})
        self.assertEqual(set(table["records"][0]["missing"]), {"b", "c"})
        self.assertEqual(set(table["records"][1]["missing"]), {"a", "b", "c"})
        allowed = self.bound_table(allow_formula=True)
        self.assertEqual(allowed["records"][1]["values"], {"a": 1.1})
        self.assertEqual(allowed["records"][1]["formula_cache_used"], ["a"])

    def test_endpoint_header_binding_is_case_sensitive_and_never_aliases(self):
        with RECOVERY.IndependentWorkbook(self.path) as workbook, self.assertRaisesRegex(ValueError, "header mismatch"):
            RECOVERY.extract_bound_table(workbook, sheet="target", header_row=1, first_data_row=2,
                id_column=1, value_columns={"a": 2}, expected_headers={1: "ID", 2: "cell a"})

    def test_exact_id_join_never_uses_suffix_case_or_favorable_mapping(self):
        left = [{"id": "a"}, {"id": "a_F"}]
        right = [{"id": "a"}, {"id": "A"}, {"id": "other"}]
        join = RECOVERY.exact_id_join(left, right)
        self.assertEqual(join["pairs"], [("a", "a")])
        self.assertEqual(join["left_only_ids"], ["a_F"])
        self.assertEqual(join["right_only_ids"], ["A", "other"])
        with self.assertRaisesRegex(ValueError, "one-to-one"):
            RECOVERY.exact_id_join(left, right, {"a": "A", "a_F": "A"})
        with self.assertRaisesRegex(ValueError, "duplicate"):
            RECOVERY.exact_id_join(left + [left[0]], right)

    def test_exact_decimal_equality_is_separate_from_format_and_fixed_relative_tolerance(self):
        left = [{"id": "x", "raw_literals": {"a": "0.10", "b": "100"}}]
        right = [{"id": "x", "raw_literals": {"a": "1e-1", "b": "100.000001"}}]
        result = RECOVERY.compare_bound_tables(left, right, ["a", "b"], absolute_tolerance=1e-8, relative_tolerance=1e-8)
        self.assertEqual(result["equal_exact_numeric_literals"], 1)
        self.assertEqual(result["equal_source_representations"], 0)
        self.assertEqual(result["exceeding_tolerance"], 0)
        absolute_only = RECOVERY.compare_bound_tables(left, right, ["a", "b"], absolute_tolerance=1e-8)
        self.assertEqual(absolute_only["exceeding_tolerance"], 1)

    def test_fixed_formula_has_hand_computed_sample_zscore_and_row_center(self):
        values = np.array([[1., 3., 5.], [2., 2., 9.], [4., 1., 6.], [7., 6., 0.]])
        result, ledger = RECOVERY.apply_fixed_equation(values, ("cell_zscore", "row_center"),
                                                      ddof=1, population_rows=[True] * 4)
        np.testing.assert_allclose(result[0], [-5 / (3 * np.sqrt(7)), 2.5 / (3 * np.sqrt(7)), 2.5 / (3 * np.sqrt(7))],
                                   rtol=0, atol=2e-16)
        np.testing.assert_allclose(result.sum(axis=1), 0, rtol=0, atol=4e-16)
        self.assertEqual(ledger["ledger"][0]["source_population_rows"], 4)

    def test_population_and_sample_ddof_factor_is_fixed_and_input_unchanged(self):
        values = np.array([[1., 3., 5.], [2., 2., 9.], [4., 1., 6.], [7., 6., 0.]])
        before = values.copy()
        a, _ = RECOVERY.apply_fixed_equation(values, ("cell_zscore",), ddof=0, population_rows=[True] * 4)
        b, _ = RECOVERY.apply_fixed_equation(values, ("cell_zscore",), ddof=1, population_rows=[True] * 4)
        np.testing.assert_allclose(b, a * np.sqrt(3 / 4), rtol=0, atol=4e-16)
        np.testing.assert_array_equal(values, before)
        shifted, _ = RECOVERY.apply_fixed_equation(values * [2, 3, 4] + [10, -20, 50],
                                                   ("cell_zscore",), ddof=1, population_rows=[True] * 4)
        np.testing.assert_allclose(shifted, b, rtol=0, atol=5e-16)

    def test_finite_cell_populations_differ_from_complete_intersection_without_imputation(self):
        values = np.array([[0., 0., 0.], [1., 3., np.nan], [4., 5., 6.]])
        complete = np.isfinite(values).all(axis=1)
        _, complete_stats = RECOVERY.column_zscores(values, np.repeat(complete[:, None], 3, axis=1), ddof=0)
        result, finite_stats = RECOVERY.column_zscores(values, np.isfinite(values), ddof=0)
        self.assertEqual([row["population_count"] for row in complete_stats["moments"]], [2, 2, 2])
        self.assertEqual([row["population_count"] for row in finite_stats["moments"]], [3, 3, 2])
        self.assertTrue(np.isnan(result[1, 2]))
        with self.assertRaisesRegex(ValueError, "contains missing"):
            RECOVERY.column_zscores(values, np.ones_like(values, dtype=bool), ddof=0)

    def test_replicate_medians_use_each_full_finite_population_then_require_all_three(self):
        values = np.array([[[0., 10., 100.]], [[2., np.nan, 104.]], [[4., 18., 108.]]])
        result, ledger = RECOVERY.replicate_activity_base(values, subtract_replicate_medians=True)
        np.testing.assert_allclose(result[[0, 2], 0], [-10 / 3, 10 / 3], rtol=0, atol=0)
        self.assertTrue(np.isnan(result[1, 0]))
        self.assertEqual([row["finite_source_count"] for row in ledger["replicate_populations"]], [3, 2, 3])
        self.assertEqual([row["median"] for row in ledger["replicate_populations"]], [2, 14, 104])
        self.assertEqual(ledger["complete_three_replicate_cells"], 2)
        self.assertFalse(ledger["imputation"])

    def test_all_twelve_variants_and_observational_equivalence_are_retained(self):
        values = np.array([[1., 3., 5.], [2., 2., 9.], [4., 1., 6.], [7., 6., 0.]])
        variants = RECOVERY.fixed_grid_for_base(values, ["a", "b", "c", "d"], "aggregate")
        self.assertEqual(len(variants), 12)
        self.assertEqual(len({row["id"] for row in variants}), 12)
        self.assertTrue(all(row["status"] == "estimable" for row in variants))
        for row in variants:
            if row["post_ddof"] is None:
                np.testing.assert_allclose(row["values"].sum(axis=1), 0, rtol=0, atol=1e-15)
        a = next(row for row in variants if row["pre_ddof"] == 0 and row["post_ddof"] == 1)
        b = next(row for row in variants if row["pre_ddof"] == 1 and row["post_ddof"] == 1)
        np.testing.assert_allclose(a["values"], b["values"], rtol=0, atol=7e-16)

    def test_undefined_grid_retains_every_nonestimable_variant(self):
        variants = RECOVERY.fixed_grid_for_base(np.ones((3, 3)), ["a", "b", "c"], "aggregate")
        self.assertEqual(len(variants), 12)
        self.assertTrue(all(row["status"] == "nonestimable" for row in variants))
        self.assertTrue(all("constant" in row["reason"] for row in variants))

    def test_single_code_test_normalizes_all_required_target_ids_once_and_preserves_cohorts(self):
        contexts = ["a", "b", "c"]
        values = np.array([[1., 2., -3.], [2., -1., -1.], [-3., 1., 2.], [7., -2., -5.]])
        ids = ["x", "y", "z", "outside"]
        source = [{"id": identifier, "values": dict(zip(contexts, row))} for identifier, row in zip(ids, values)]
        expected = (values[:3] - values[:3].mean(axis=0)) / values[:3].std(axis=0, ddof=1)
        target = [{"id": identifier, "values": dict(zip(contexts, row))} for identifier, row in zip(ids[:3], expected)]
        result = RECOVERY.published_ess_subset_test(source, target, contexts, old_ids={"x", "y"}, old_admitted_ids={"x"})
        self.assertEqual(result["status"], "estimable")
        self.assertTrue(all(row["all_within_tolerance"] for row in result["comparisons"]))
        self.assertEqual([row["id_count"] for row in result["comparisons"]], [3, 2, 1])
        self.assertTrue(all(not row["normalization_recomputed_on_cohort"] for row in result["comparisons"]))
        broken = [*target, {"id": "unmatched", "values": dict(zip(contexts, [0., 0., 0.]))}]
        failed = RECOVERY.published_ess_subset_test(source, broken, contexts, old_ids=set(), old_admitted_ids=set())
        self.assertEqual(failed["status"], "nonestimable")
        self.assertEqual(failed["missing_ids"], ["unmatched"])

    def test_single_code_test_rejects_any_required_incomplete_or_duplicate_measurement(self):
        source = [{"id": "a", "values": {"x": 1, "y": 2, "z": 3}},
                  {"id": "b", "values": {"x": 2, "y": 3, "z": 4}}]
        target = [{"id": "a", "values": {"x": 0, "y": 0, "z": 0}},
                  {"id": "b", "values": {"x": 0, "y": 0}}]
        result = RECOVERY.published_ess_subset_test(source, target, ["x", "y", "z"], old_ids=set(), old_admitted_ids=set())
        self.assertEqual(result["status"], "nonestimable")
        duplicated = RECOVERY.published_ess_subset_test(source + [source[0]], target, ["x", "y", "z"],
                                                       old_ids=set(), old_admitted_ids=set())
        self.assertEqual(duplicated["status"], "nonestimable")
        self.assertIn("duplicate", duplicated["reason"])

    def test_source_moments_keep_zero_one_and_many_observations_distinct(self):
        values = np.array([[1., np.nan, 3.], [3., np.nan, np.nan]])
        result = RECOVERY.matrix_moments(values, ["a", "b", "c"])
        self.assertEqual([row["finite_count"] for row in result["by_context"]], [2, 0, 1])
        self.assertEqual(result["by_context"][0]["mean"], 2)
        self.assertEqual(result["by_context"][0]["population_sd"], 1)
        self.assertAlmostEqual(result["by_context"][0]["sample_sd"], np.sqrt(2))
        self.assertIsNone(result["by_context"][1]["mean"])
        self.assertIsNone(result["by_context"][2]["sample_sd"])
        self.assertEqual(result["complete_context_rows"], 0)
        self.assertIsNone(result["complete_row_sum"]["maximum_absolute"])
        json.dumps(result, allow_nan=False)

    def test_numeric_or_formula_header_artifacts_are_redacted_before_semantic_binding(self):
        numeric = RECOVERY.SourceCell(1, 8, "n", 12345.6, "n", "12345.6", False, None)
        formula = RECOVERY.SourceCell(1, 9, "f", "=A2", "n", "9876.5", True, "A2")
        for cell in [numeric, formula]:
            report = cell.header_json()
            self.assertIsNone(report["value"])
            self.assertTrue(report["nontext_header_value_redacted"])
            serialized = json.dumps(report)
            self.assertNotIn("12345.6", serialized)
            self.assertNotIn("9876.5", serialized)
            self.assertNotIn("A2", serialized)

    def test_manifest_verification_checks_both_plain_and_sized_sha256_entries(self):
        folder = self.root / "bundle"
        folder.mkdir()
        (folder / "data").write_bytes(b"exact")
        digest = hashlib.sha256(b"exact").hexdigest()
        for entry in [digest, {"bytes": 5, "sha256": digest}]:
            (folder / "manifest.json").write_text(json.dumps({"files": {"data": entry}}))
            RECOVERY.verified_manifest(folder, "manifest.json")
        (folder / "manifest.json").write_text(json.dumps({"files": {"data": {"bytes": 6, "sha256": digest}}}))
        with self.assertRaisesRegex(ValueError, "integrity mismatch"):
            RECOVERY.verified_manifest(folder, "manifest.json")

    def test_exact_constant_fractional_column_is_not_rescued_by_roundoff_sd(self):
        matrix = np.full((1000, 3), 0.1)
        with self.assertRaisesRegex(ValueError, "constant"):
            RECOVERY.column_zscores(matrix, np.ones_like(matrix, dtype=bool), ddof=1)

    def test_exact_ESS_cohort_does_not_rescue_missing_or_nonfinite_required_ids(self):
        ids = ["a", "b", "c"]
        values = np.array([[1., 2., 3.], [2., 4., 1.], [4., 1., 2.]])
        for ddof in [0, 1]:
            missing = RECOVERY.exact_ess_cohort_route(ids, values, ["a", "b", "required"], ddof=ddof)
            self.assertEqual(missing["status"], "nonestimable")
            self.assertEqual(missing["missing_ids"], ["required"])
            broken = values.copy()
            broken[1, 1] = np.nan
            nonfinite = RECOVERY.exact_ess_cohort_route(ids, broken, ids, ddof=ddof)
            self.assertEqual(nonfinite["status"], "nonestimable")
            self.assertEqual(nonfinite["nonfinite_ids"], ["b"])

    def test_fixed_ESS_population_excludes_other_rows_without_estimation_on_old_intersections(self):
        ids = ["a", "b", "c", "not-in-ESS"]
        values = np.array([[1., 2., 3.], [2., 4., 1.], [4., 1., 2.], [1000., -999., 777.]])
        result = RECOVERY.exact_ess_cohort_route(ids, values, ids[:3], ddof=1)
        self.assertEqual(result["status"], "estimable")
        expected = (values[:3] - values[:3].mean(axis=0)) / values[:3].std(axis=0, ddof=1)
        expected -= expected.mean(axis=1, keepdims=True)
        np.testing.assert_array_equal(result["values"], expected)
        self.assertEqual([item["population_count"] for item in result["normalization"]["moments"]], [3, 3, 3])


if __name__ == "__main__":
    unittest.main()
